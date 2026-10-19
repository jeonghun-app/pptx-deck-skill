#!/usr/bin/env python3
"""Regression test for deck-craft-lint.py using a synthetic deck (python-pptx)."""
import json, subprocess, sys, tempfile
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.util import Pt, Emu

HERE = Path(__file__).resolve().parent
prs = Presentation(); prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
blank = prs.slide_layouts[6]

def box(slide, x, y, w, h, text, size, color=None, fill=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill)
    else: s.fill.background()
    r = s.text_frame.paragraphs[0].add_run(); r.text = text; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    return s

# slide 1: clean
s1 = prs.slides.add_slide(blank)
box(s1, 72, 48, 800, 33, "월 비용이 42% 줄어듭니다", 25.5, "1F2328")
box(s1, 72, 150, 400, 40, "통합 후 비용 비중 42% 감소", 12.75, "1F2328")
# slide 2: defects
s2 = prs.slides.add_slide(blank)
box(s2, 72, 48, 800, 33, "처리량이 7배 늘었습니다", 25.5, "1F2328")
box(s2, 300, -40, 400, 24, "복사해서 사용하세요", 10.5, "1F2328")           # off-canvas text
box(s2, 72, 150, 400, 40, "메시지를 기술하세요", 12.75, "1F2328")           # template wording
box(s2, 72, 220, 200, 30, "8pt value 58", 8, "1F2328")                      # too small
box(s2, 72, 270, 200, 30, "White on teal", 14, "FFFFFF", "4FB3BF")          # low contrast
o = s2.shapes.add_shape(MSO_SHAPE.OVAL, Pt(500), Pt(300), Pt(40), Pt(38))   # distorted circle
cd = CategoryChartData(); cd.categories = ["a", "b"]; cd.add_series("s", (45, 55))
ch = s2.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED_100, Pt(600), Pt(150), Pt(300), Pt(200), cd).chart
pt = ch.plots[0].series[0].points[0]; pt.data_label.text_frame.text = "41"     # typed numeric label

# slide 3: legitimate counterexamples that must not fail (and must not warn on these rules)
s3 = prs.slides.add_slide(blank)
import io
from PIL import Image
buf = io.BytesIO(); Image.new("RGB", (64, 36), (200, 210, 230)).save(buf, "PNG"); buf.seek(0)
s3.shapes.add_picture(buf, Pt(-20), Pt(-20), Pt(1000), Pt(580))                     # intended full bleed
box(s3, 72, 48, 800, 33, "권장 구성의 비용이 가장 낮습니다", 25.5, "1F2328")
tb = s3.shapes.add_table(3, 2, Pt(72), Pt(150), Pt(500), Pt(90)).table             # quiet table: 9pt header
for c in range(2):
    cell = tb.cell(0, c); cell.text = "구분"; cell.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    cell = tb.cell(1, c); cell.text = "설명"; cell.text_frame.paragraphs[0].runs[0].font.size = Pt(9.75)
s3.shapes.add_shape(MSO_SHAPE.OVAL, Pt(600), Pt(300), Pt(64), Pt(40))               # intended ellipse
box(s3, 600, 360, 200, 40, "권장", 18, "FFFFFF", "2F5BD3")                          # white on accent, ok
cd3 = CategoryChartData(); cd3.categories = ["a", "b"]; cd3.add_series("s", (45, 55))
ch3 = s3.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Pt(600), Pt(150), Pt(300), Pt(120), cd3).chart
ch3.plots[0].series[0].points[0].data_label.text_frame.text = "45"                   # typed but equal to data
box(s3, 72, 506, 400, 14, "푸터 · 1 / 3", 9.75, "5E6672")                           # footer chrome at the minimum

with tempfile.TemporaryDirectory() as d:
    f = Path(d) / "t.pptx"; prs.save(f); out = Path(d) / "r.json"
    rc = subprocess.run([sys.executable, str(HERE / "deck-craft-lint.py"), str(f), "--output", str(out)]).returncode
    r = json.loads(out.read_text())
    rules = {(x["slide"], x["rule"], x["severity"]) for x in r["findings"]}
    expect = [(2, "off-canvas", "fail"), (2, "template-text", "fail"), (2, "min-size", "fail"),
              (2, "light-text-on-fill", "fail"), (2, "distorted-circle", "warn"), (2, "chart-typed-label", "fail"),
              (2, "headline-number", "warn")]
    missing = [e for e in expect if e not in rules]
    clean = [x for x in r["findings"] if x["slide"] == 1 and x["severity"] == "fail"]
    assert rc == 1, rc
    assert not missing, f"missing findings: {missing}\n{json.dumps(r['findings'], ensure_ascii=False, indent=1)}"
    assert not clean, clean
    legit = [x for x in r["findings"] if x["slide"] == 3 and (x["severity"] == "fail" or
             (x["severity"] == "warn" and x["rule"] in {"off-canvas", "distorted-circle", "chart-typed-label", "light-text-on-fill", "min-size", "footer-clearance"}))]
    assert not legit, f"false positives on legitimate counterexamples: {legit}"

# Brand profile: thresholds, palette and fonts come from brand.yaml.
BRAND = """
canvas: {width: 960, height: 540}
type: {min: 10}
layout: {body_bottom: 466}
colors: {ink: 1F2328, accent: 2F5BD3, positive: 1A7F4B, series: [2F5BD3, 7A8699]}
fonts:
  heading: {latin: Arial, east_asian: Malgun Gothic}
  body: {latin: Arial, east_asian: Malgun Gothic}
"""
bp = Presentation(); bp.slide_width, bp.slide_height = Emu(12192000), Emu(6858000)
b1 = bp.slides.add_slide(bp.slide_layouts[6])
r = box(b1, 72, 48, 800, 33, "On-brand headline", 18, "1F2328").text_frame.paragraphs[0].runs[0]; r.font.name = "Arial"
r = box(b1, 72, 150, 400, 40, "Off-brand color and font", 14, "C2185B").text_frame.paragraphs[0].runs[0]; r.font.name = "Comic Sans MS"
box(b1, 72, 220, 400, 30, "9.9pt text", 9.9, "1F2328")   # passes the neutral 9.75pt floor, fails brand min 10
box(b1, 72, 270, 12, 12, "", 10, fill="7894E2")            # chart key swatch: an accent ramp tint is on-brand

with tempfile.TemporaryDirectory() as d:
    f = Path(d) / "b.pptx"; bp.save(f); y = Path(d) / "brand.yaml"; y.write_text(BRAND)
    def run(*extra):
        out = Path(d) / "o.json"
        rc = subprocess.run([sys.executable, str(HERE / "deck-craft-lint.py"), str(f), "--output", str(out), *extra]).returncode
        return rc, json.loads(out.read_text())
    rc0, plain = run()
    rc1, branded = run("--brand", str(y))
    got0 = {(x["rule"], x["severity"]) for x in plain["findings"]}
    got1 = {(x["rule"], x["severity"]) for x in branded["findings"]}
    assert rc0 == 0 and ("min-size", "fail") not in got0, plain["findings"]
    assert not any(r_ in {"off-brand-color", "off-brand-font"} for r_, _ in got0), "palette checks need --brand"
    assert rc1 == 1 and ("min-size", "fail") in got1, branded["findings"]
    assert ("off-brand-color", "warn") in got1 and ("off-brand-font", "warn") in got1, branded["findings"]
    assert branded["thresholds"]["min_pt"] == 10 and branded["thresholds"]["palette_checked"]
    assert all("7894E2" not in x["evidence"] for x in branded["findings"]), branded["findings"]
    assert all("1F2328" not in x["evidence"] and "Arial" not in x["evidence"]
               for x in branded["findings"] if x["rule"] in {"off-brand-color", "off-brand-font"})
    assert subprocess.run([sys.executable, str(HERE / "deck-craft-lint.py"), str(f), "--brand", str(Path(d) / "none.yaml")],
                          capture_output=True).returncode == 2

    # Without PyYAML the brand profile is ignored gracefully, with a recorded limitation.
    import importlib.util
    spec = importlib.util.spec_from_file_location("deck_craft_lint", HERE / "deck-craft-lint.py")
    mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod)
    saved = sys.modules.get("yaml"); sys.modules["yaml"] = None
    import contextlib, io as _io
    try:
        with contextlib.redirect_stderr(_io.StringIO()):  # the fallback warning is expected here
            cfg, notes = mod.load_brand(str(y))
    finally:
        if saved is None: sys.modules.pop("yaml", None)
        else: sys.modules["yaml"] = saved
    assert cfg["palette"] is None and cfg["min_pt"] == mod.DEFAULTS["min_pt"] and notes, (cfg, notes)
print("deck-craft-lint regression: OK")
