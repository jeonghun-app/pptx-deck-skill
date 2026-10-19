"""Regression tests for cross-track behaviour: measurement, big titles, bullet lists."""
import shutil
import sys
from pathlib import Path

import pytest
import yaml
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR

SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))

from deckkit import build, load_brand  # noqa: E402
from deckkit import measure  # noqa: E402
from lint_deck import lint  # noqa: E402

CONTRACT = {"audience": "a", "objective": "o", "desired_action": "d"}


def deck(tmp_path, slides, brand=None):
    p = tmp_path / "d.deck.yaml"
    p.write_text(yaml.safe_dump({"meta": CONTRACT, "slides": slides}, allow_unicode=True), encoding="utf-8")
    out = tmp_path / "d.pptx"
    report = build(p, out, brand)
    return Presentation(str(out)), report


def test_heuristic_fallback_without_fonts():
    measure._FILES.clear()
    assert measure.text_width("abc", 10) == pytest.approx(sum(measure.char_em(c) for c in "abc") * 10)


@pytest.mark.skipif(not shutil.which("fc-match"), reason="fontconfig not available")
def test_real_font_measurement_tracks_the_face(tmp_path):
    narrow, wide = tmp_path / "narrow.yaml", tmp_path / "wide.yaml"
    narrow.write_text("fonts: {body: {latin: Liberation Serif}}\n", encoding="utf-8")
    wide.write_text("fonts: {body: {latin: DejaVu Sans}}\n", encoding="utf-8")
    text = "Most blockers surface only after the branch is cut"
    measure.set_fonts(load_brand(narrow))
    w_narrow = measure.text_width(text, 20)
    measure.set_fonts(load_brand(wide))
    w_wide = measure.text_width(text, 20)
    assert w_wide > w_narrow * 1.1
    measure.set_fonts(load_brand())


def test_brand_font_file_overrides_fontconfig(tmp_path):
    match = measure._match("DejaVu Sans")
    if not match:
        pytest.skip("fontconfig not available")
    brand = tmp_path / "b.yaml"
    brand.write_text(f"fonts: {{body: {{latin: Not A Real Face, file: '{match}'}}}}\n", encoding="utf-8")
    measure.set_fonts(load_brand(brand))
    assert measure._FILES["body"][0] == match
    measure.set_fonts(load_brand())


def test_big_titles_grow_upward(tmp_path):
    prs, _ = deck(tmp_path, [
        {"layout": "cover", "title": "A cover title", "subtitle": "Sub", "notes": "n"},
        {"layout": "statement", "text": "One sentence to keep", "notes": "n"}])
    for slide, name in ((prs.slides[0], "Title"), (prs.slides[1], "Statement")):
        box = next(s for s in slide.shapes if s.name == name)
        assert box.text_frame.vertical_anchor == MSO_ANCHOR.BOTTOM


def test_bullets_are_one_text_box_with_paragraphs(tmp_path):
    prs, report = deck(tmp_path, [{
        "layout": "bullets", "title": "Three points carry the decision today", "notes": "n",
        "bullets": ["First point", {"text": "Second point", "sub": ["detail a", "detail b"]}, "Third"]}])
    boxes = [s for s in prs.slides[0].shapes if s.name == "Bullets"]
    assert len(boxes) == 1
    assert [p.text for p in boxes[0].text_frame.paragraphs] == [
        "First point", "Second point", "detail a", "detail b", "Third"]
    assert not [w for w in report["warnings"] if w["level"] == "ERROR"]


def test_summary_expected_only_in_sectioned_decks(tmp_path):
    body = [{"layout": "bullets", "title": f"Finding {k} changes the plan for us", "bullets": ["a"],
             "notes": "n"} for k in range(10)]
    p = tmp_path / "c.deck.yaml"
    p.write_text(yaml.safe_dump({"meta": CONTRACT, "slides": body}), encoding="utf-8")
    assert not any("no summary" in f["message"] for f in lint(p))


def test_matrix_commentary_is_supported(tmp_path):
    spec = {"layout": "matrix", "title": "Two options combine value with easy adoption", "notes": "n",
            "x": {"label": "Ease", "low": "Hard", "high": "Easy"},
            "y": {"label": "Value", "low": "Low", "high": "High"},
            "quadrants": [{"name": n} for n in ("A", "B", "C", "D")],
            "commentary": {"text": "Start with B."}}
    p = tmp_path / "m.deck.yaml"
    p.write_text(yaml.safe_dump({"meta": CONTRACT, "slides": [spec]}), encoding="utf-8")
    assert not any("commentary has no effect" in f["message"] for f in lint(p))
