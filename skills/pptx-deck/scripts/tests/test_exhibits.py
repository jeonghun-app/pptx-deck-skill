"""Exhibit features: python -m pytest skills/pptx-deck/scripts/tests/test_exhibits.py"""
import re
import subprocess
import sys
from pathlib import Path

import pytest
import yaml
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

SCRIPTS = Path(__file__).resolve().parents[1]
SKILL = SCRIPTS.parent
sys.path.insert(0, str(SCRIPTS))

from deckkit import build, load_brand  # noqa: E402
from deckkit import exhibits  # noqa: E402
from lint_deck import lint  # noqa: E402

META = {"audience": "a", "objective": "o", "desired_action": "d"}
BRAND = load_brand()
GRID = 12


def make(tmp_path, slides, meta=None, brand=None, name="d"):
    deck = tmp_path / f"{name}.deck.yaml"
    deck.write_text(yaml.safe_dump({"meta": {**META, **(meta or {})}, "slides": slides},
                                   allow_unicode=True, sort_keys=False), encoding="utf-8")
    report = build(deck, tmp_path / f"{name}.pptx", brand)
    return Presentation(str(tmp_path / f"{name}.pptx")), report, deck


def chart_slide(chart, **extra):
    return {"layout": "chart", "title": "A headline that states the result", "chart": chart,
            "source": "Illustrative", "notes": "n", **extra}


def shapes(slide, name):
    return [s for s in slide.shapes if s.name == name]


def pt(emu):
    return Emu(emu).pt


def charts(slide):
    return [s for s in slide.shapes if getattr(s, "has_chart", False) and s.has_chart]


def series_fills(frame):
    """Fill color of each series, in order (None when the series has no solid fill)."""
    out = []
    for ser in frame.chart._chartSpace.iter(qn("c:ser")):
        sp = ser.find(qn("c:spPr"))
        clr = sp.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}") if sp is not None else None
        line = sp.find(f"{qn('a:ln')}/{qn('a:solidFill')}/{qn('a:srgbClr')}") if sp is not None else None
        out.append((clr if clr is not None else line).get("val") if (clr is not None or line is not None) else None)
    return out


def point_fills(frame, series=0):
    ser = list(frame.chart._chartSpace.iter(qn("c:ser")))[series]
    return {int(d.find(qn("c:idx")).get("val")):
            d.find(f".//{qn('a:solidFill')}/{qn('a:srgbClr')}").get("val") for d in ser.findall(qn("c:dPt"))}


def manual_layout(frame):
    ml = frame.chart._chartSpace.find(f".//{qn('c:plotArea')}/{qn('c:layout')}/{qn('c:manualLayout')}")
    if ml is None:
        return None
    return {k: float(ml.find(qn(f"c:{k}")).get("val")) for k in ("x", "y", "w", "h")}


def axis_scale(frame):
    va = frame.chart.value_axis
    return va.minimum_scale, va.maximum_scale


def plot_rect(frame):
    """The inner plot rectangle in slide pt, read back from the written layout."""
    m = manual_layout(frame)
    x, y, w, h = (pt(v) for v in (frame.left, frame.top, frame.width, frame.height))
    return x + m["x"] * w, y + m["y"] * h, m["w"] * w, m["h"] * h


# -- 1. footer --------------------------------------------------------------

def test_string_source_renders_as_before(tmp_path):
    prs, _, _ = make(tmp_path, [chart_slide({"categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]})])
    (src,) = shapes(prs.slides[0], "Source")
    assert src.text_frame.text == "Source: Illustrative"
    bottom = float(BRAND["layout.body_bottom"])
    assert pt(src.top) == pytest.approx(bottom - BRAND["type.caption"] * 1.5, abs=0.01)
    assert pt(src.height) == pytest.approx(BRAND["type.caption"] * 1.5, abs=0.01)


def test_footer_object_orders_notes_meta_source_and_grows_upward(tmp_path):
    source = {"notes": ["Top two boxes.", "Excludes trials."], "question": "Would you renew?",
              "base": "Active members", "n": {"Total": 1500, "Segment A": 136}, "as_of": "Sep 2026",
              "rounding": True, "text": "Illustrative survey; own analysis"}
    prs, report, _ = make(tmp_path, [chart_slide({"categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]},
                                                 source=source)])
    (src,) = shapes(prs.slides[0], "Source")
    paras = [p.text for p in src.text_frame.paragraphs]
    assert paras[0].startswith("1) Top two boxes.  2) Excludes trials.") and "rounding" in paras[0]
    assert paras[1] == ("Q: Would you renew? · Base: Active members · n = 1,500 (Total), 136 (Segment A)"
                        " · As of Sep 2026")
    assert paras[2] == "Source: Illustrative survey; own analysis"
    assert pt(src.top + src.height) == pytest.approx(float(BRAND["layout.body_bottom"]), abs=0.01)
    assert pt(src.height) > BRAND["type.caption"] * 1.5 * 2.5  # three lines, grown upward
    (frame,) = charts(prs.slides[0])  # the body limit shrinks with it
    assert pt(frame.top + frame.height) == pytest.approx(pt(src.top) - BRAND["layout.gutter"] / 2, abs=0.01)
    assert not [w for w in report["warnings"] if w["level"] == "ERROR"]


def test_footer_uses_korean_labels(tmp_path):
    source = {"question": "갱신하시겠습니까?", "base": "활성 회원", "n": 300, "as_of": "2026년 9월",
              "text": "예시 조사"}
    prs, _, _ = make(tmp_path, [chart_slide({"categories": ["가", "나"], "series": [{"name": "s", "values": [1, 2]}]},
                                            title="북부 회원의 갱신 의향이 가장 높다", source=source)],
                     meta={"language": "ko"})
    (src,) = shapes(prs.slides[0], "Source")
    assert [p.text for p in src.text_frame.paragraphs] == [
        "질문: 갱신하시겠습니까? · 대상: 활성 회원 · n = 300 · 2026년 9월 기준", "출처: 예시 조사"]


def test_footer_accepts_one_note_as_a_string(tmp_path):
    prs, _, _ = make(tmp_path, [chart_slide({"categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]},
                                            source={"notes": "Fiscal years.", "text": "Ledger"})])
    (src,) = shapes(prs.slides[0], "Source")
    assert [p.text for p in src.text_frame.paragraphs] == ["1) Fiscal years.", "Source: Ledger"]


def test_brand_overrides_footer_wording(tmp_path):
    brand = tmp_path / "brand.yaml"
    brand.write_text("copy: {source_label: Data, exhibit_labels: {base: 'Sample: {}'}}\n", encoding="utf-8")
    prs, _, _ = make(tmp_path, [chart_slide({"categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]},
                                            source={"base": "All", "text": "Ledger"})], brand=brand)
    (src,) = shapes(prs.slides[0], "Source")
    assert [p.text for p in src.text_frame.paragraphs] == ["Sample: All", "Data: Ledger"]


def test_lint_footer_rules(tmp_path):
    long_note = "A definition that keeps going for a while to fill the line " * 3
    slides = [
        chart_slide({"categories": ["a", "b"], "series": [{"name": "s", "values": [0.4, 0.5]}]}, survey=True,
                    source={"text": "Survey"}),
        chart_slide({"categories": ["a", "b"], "series": [{"name": "s", "values": [0.4, 0.5]}]}, survey=True,
                    source={"n": 900, "notes": [long_note, long_note], "question": "Q?", "text": "Survey"}),
        chart_slide({"categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]},
                    source={"notes": ["No source text"]}),
    ]
    _, _, deck = make(tmp_path, slides)
    found = [(f["slide"], f["message"]) for f in lint(deck)]
    assert any(s == 1 and "without a sample size" in m for s, m in found)
    assert not any(s == 2 and "without a sample size" in m for s, m in found)
    assert any(s == 2 and "footer runs to" in m for s, m in found)
    assert any(s == 3 and "without a source line" in m for s, m in found)


# -- 2/3. highlight and series colors ------------------------------------------

def test_highlight_accepts_a_list_and_category_names(tmp_path):
    prs, _, _ = make(tmp_path, [chart_slide({"type": "column", "categories": ["A", "B", "C", "D"],
                                             "series": [{"name": "s", "values": [1, 2, 3, 4]}],
                                             "highlight": [0, "C"]})])
    (frame,) = charts(prs.slides[0])
    accent = BRAND.hex("accent")
    context = BRAND.tint("muted", 0.45)
    assert point_fills(frame) == {0: accent, 2: accent}
    assert series_fills(frame) == [context]


def test_highlight_series_puts_focus_in_accent_and_context_in_grey(tmp_path):
    chart = {"type": "column", "categories": ["Q1", "Q2"], "highlight_series": 1,
             "series": [{"name": n, "values": [1, 2]} for n in ("A", "B", "C")]}
    prs, _, _ = make(tmp_path, [chart_slide(chart)])
    (frame,) = charts(prs.slides[0])
    assert series_fills(frame) == [BRAND.tint("muted", 0.45), BRAND.hex("accent"), BRAND.tint("muted", 0.75)]


def test_ordered_series_get_an_accent_ramp_and_nominal_the_palette(tmp_path):
    base = {"type": "column", "categories": ["Q1", "Q2"],
            "series": [{"name": n, "values": [1, 2]} for n in ("A", "B", "C")]}
    prs, _, _ = make(tmp_path, [chart_slide(base), chart_slide({**base, "series_kind": "nominal"})])
    ordered, nominal = (charts(s)[0] for s in prs.slides)
    assert series_fills(ordered) == [BRAND.tint("accent", t) for t in (0, 0.35, 0.7)]
    assert series_fills(nominal) == [BRAND.series(i) for i in range(3)]
    brand = tmp_path / "brand.yaml"
    brand.write_text("colors: {series_mode: palette}\n", encoding="utf-8")
    prs, _, _ = make(tmp_path, [chart_slide(base)], brand=brand, name="p")
    assert series_fills(charts(prs.slides[0])[0]) == [BRAND.series(i) for i in range(3)]


def test_single_series_keeps_the_first_series_color(tmp_path):
    prs, _, _ = make(tmp_path, [chart_slide({"categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]})])
    assert series_fills(charts(prs.slides[0])[0]) == [BRAND.series(0)]


def test_lint_highlight_rules(tmp_path):
    two = [{"name": "A", "values": [1, 2]}, {"name": "B", "values": [2, 1]}]
    _, _, deck = make(tmp_path, [
        chart_slide({"categories": ["x", "y"], "series": two, "highlight": 0}),
        chart_slide({"categories": ["x", "y"], "series": two[:1], "highlight": "z"}),
        chart_slide({"categories": ["x", "y"], "series": two, "highlight_series": 5}),
    ])
    found = [(f["level"], f["slide"], f["message"]) for f in lint(deck)]
    assert ("WARN", 1, "highlight marks categories of a single series; use highlight_series") in found
    assert any(lv == "ERROR" and s == 2 and "'z' is not a category" in m for lv, s, m in found)
    assert any(lv == "ERROR" and s == 3 and "highlight_series 5" in m for lv, s, m in found)


# -- 4. line end labels --------------------------------------------------------

def test_line_end_labels_replace_the_legend(tmp_path):
    chart = {"type": "line", "categories": ["2024", "2025", "2026"],
             "series": [{"name": "North", "values": [10, 20, 30]}, {"name": "East", "values": [9, 15, 29]},
                        {"name": "West", "values": [8, 12, 28]}]}
    prs, report, _ = make(tmp_path, [chart_slide(chart)])
    slide = prs.slides[0]
    (frame,) = charts(slide)
    assert not frame.chart.has_legend
    assert manual_layout(frame) is not None and axis_scale(frame) == (0, 30)
    labels = shapes(slide, "End label")
    assert sorted(lb.text_frame.text for lb in labels) == ["East  29", "North  30", "West  28"]
    boxes = sorted((pt(lb.top), pt(lb.height)) for lb in labels)
    for (top_a, h_a), (top_b, _) in zip(boxes, boxes[1:]):
        assert top_b >= top_a + h_a - 0.01  # close values are nudged apart
    px, _, pw, _ = plot_rect(frame)
    assert all(pt(lb.left) >= px + pw for lb in labels)  # right of the plot, in the padding
    assert not report["warnings"]


def test_many_lines_keep_the_legend_and_lint_warns(tmp_path):
    chart = {"type": "line", "categories": ["a", "b"],
             "series": [{"name": f"S{i}", "values": [i, i + 1]} for i in range(7)]}
    prs, _, deck = make(tmp_path, [chart_slide(chart)])
    (frame,) = charts(prs.slides[0])
    assert frame.chart.has_legend and not shapes(prs.slides[0], "End label")
    assert any("7 lines with no highlight" in f["message"] for f in lint(deck))


# -- 5. waterfall ----------------------------------------------------------------

def test_waterfall_bridges_start_to_end(tmp_path):
    chart = {"type": "waterfall", "start": {"label": "2025", "value": 120},
             "steps": [{"label": "New", "value": 34}, {"label": "Churn", "value": -27}], "total": "2026"}
    prs, report, _ = make(tmp_path, [chart_slide(chart)])
    (frame,) = charts(prs.slides[0])
    plot = frame.chart.plots[0]
    assert list(plot.categories) == ["2025", "New", "Churn", "2026"]
    values = {s.name: list(s.values) for s in plot.series}
    assert values["Base"] == [0, 120, 127, 0]
    assert values["Total"][0] == 120 and values["Total"][3] == 127  # end = start + sum(steps)
    assert values["Increase"][1] == 34 and values["Decrease"][2] == 27
    fills = series_fills(frame)
    assert fills[0] is None  # invisible base
    assert fills[1:] == [BRAND.hex("ink"), BRAND.hex("positive"), BRAND.hex("negative")]
    labels = [s.text_frame.text for s in shapes(prs.slides[0], "Value label")]
    assert labels == ["120", "+34", "−27", "127"]
    assert not frame.chart.has_legend and not report["warnings"]


def test_unsupported_annotations_and_estimates_are_refused(tmp_path):
    pie = {"type": "pie", "categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}],
           "annotations": [{"kind": "ref_line", "value": 1}]}
    fall = {"type": "waterfall", "steps": [{"label": "a", "value": 2}], "estimates": {"from": 0},
            "total": "End"}
    mean = {"categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}],
            "annotations": [{"kind": "ref_line", "value": "mean", "series": 3}]}
    _, report, deck = make(tmp_path, [chart_slide(pie), chart_slide(fall), chart_slide(mean)])
    assert any(w["slide"] == 1 and "ignored on 'pie'" in w["message"] for w in report["warnings"])
    found = [(f["level"], f["slide"], f["message"]) for f in lint(deck)]
    assert any(lv == "ERROR" and s == 2 and "estimates work on" in m for lv, s, m in found)
    assert any(lv == "ERROR" and s == 3 and "ref_line value" in m for lv, s, m in found)


def test_waterfall_below_zero_is_an_error(tmp_path):
    chart = {"type": "waterfall", "start": {"label": "a", "value": 5}, "steps": [{"label": "b", "value": -9}]}
    _, report, deck = make(tmp_path, [chart_slide(chart)])
    assert any(w["level"] == "ERROR" and "below zero" in w["message"] for w in report["warnings"])
    assert any(f["level"] == "ERROR" and "below zero" in f["message"] for f in lint(deck))


# -- 6. estimates ------------------------------------------------------------------

def test_line_estimates_are_a_dashed_continuation(tmp_path):
    chart = {"type": "line", "categories": ["2024", "2025", "2026", "2027"],
             "series": [{"name": "A", "values": [1, 2, 3, 4]}, {"name": "B", "values": [2, 2, 2, 2]}],
             "estimates": {"from": "2026", "suffix": "e"}, "end_labels": False}
    prs, _, _ = make(tmp_path, [chart_slide(chart)])
    (frame,) = charts(prs.slides[0])
    plot = frame.chart.plots[0]
    assert list(plot.categories) == ["2024", "2025", "2026e", "2027e"]
    names = [s.name for s in plot.series]
    assert names == ["A", "B", "A (estimate)", "B (estimate)"]
    assert list(plot.series[0].values) == [1, 2, None, None]
    assert list(plot.series[2].values) == [None, 2, 3, 4]  # starts at the last actual point
    xml = frame.chart._chartSpace.xml
    assert xml.count('<a:prstDash val="dash"/>') == 2
    hidden = {int(e.find(qn("c:idx")).get("val")) for e in frame.chart._chartSpace.iter(qn("c:legendEntry"))}
    assert hidden == {2, 3}  # the legend lists each series once
    (src,) = shapes(prs.slides[0], "Source")
    assert [p.text for p in src.text_frame.paragraphs] == ["e = estimate.", "Source: Illustrative"]


def test_bar_estimates_are_lighter_and_noted(tmp_path):
    chart = {"type": "column", "categories": ["2024", "2025", "2026"],
             "series": [{"name": "A", "values": [1, 2, 3]}], "estimates": {"from": 2}}
    prs, _, _ = make(tmp_path, [chart_slide(chart)])
    (frame,) = charts(prs.slides[0])
    assert point_fills(frame) == {2: BRAND.tint(BRAND.series(0), 0.55)}
    (src,) = shapes(prs.slides[0], "Source")
    assert src.text_frame.paragraphs[0].text == "Values from 2026 on are estimates."


# -- 7. annotations ------------------------------------------------------------------

def test_annotations_land_on_the_data(tmp_path):
    chart = {"type": "column", "categories": ["A", "B", "C"], "number_format": "0%",
             "series": [{"name": "s", "values": [0.42, 0.51, 0.60]}],
             "annotations": [{"kind": "ref_line", "value": 0.5, "label": "Target {value}"},
                             {"kind": "delta", "from": "A", "to": "C", "fmt": "pp"},
                             {"kind": "callout", "at": "B", "text": "Now {value}"},
                             {"kind": "band", "from": "B", "to": "C", "label": "Pilot"}]}
    prs, report, _ = make(tmp_path, [chart_slide(chart)])
    slide = prs.slides[0]
    (frame,) = charts(slide)
    lo, hi = axis_scale(frame)
    x, y, w, h = plot_rect(frame)
    (ref,) = shapes(slide, "Reference line")
    assert pt(ref.top) == pytest.approx(y + h * (1 - (0.5 - lo) / (hi - lo)), abs=0.05)
    assert pt(ref.left) == pytest.approx(x, abs=0.05) and pt(ref.width) == pytest.approx(w, abs=0.05)
    assert shapes(slide, "Reference label")[0].text_frame.text == "Target 50%"
    assert shapes(slide, "Delta")[0].text_frame.text == "+18pp"  # computed, not typed
    assert shapes(slide, "Callout")[0].text_frame.text == "Now 51%"
    (band,) = shapes(slide, "Band")
    assert pt(band.left) == pytest.approx(x + w / 3, abs=0.05) and pt(band.width) == pytest.approx(w * 2 / 3, abs=0.05)
    order = [s.name for s in slide.shapes]
    assert order.index("Band") < order.index("Chart") < order.index("Delta")  # band behind the data
    pills = shapes(slide, "Delta") + shapes(slide, "Callout")
    a, b = ((pt(p.left), pt(p.top), pt(p.left + p.width), pt(p.top + p.height)) for p in pills)
    assert a[2] <= b[0] or b[2] <= a[0] or a[3] <= b[1] or b[3] <= a[1]  # pills never overlap
    assert all(pt(p.top + p.height) <= y for p in pills)  # lanes sit above the plot
    assert not report["warnings"]


def test_deltas_sharing_an_end_get_separate_lanes(tmp_path):
    chart = {"type": "column", "categories": ["A", "B", "C"], "series": [{"name": "s", "values": [1, 2, 3]}],
             "annotations": [{"kind": "delta", "from": "A", "to": "B", "fmt": "x"},
                             {"kind": "delta", "from": "B", "to": "C", "fmt": "pct"},
                             {"kind": "callout", "at": "B", "text": "Launch"}]}
    prs, _, _ = make(tmp_path, [chart_slide(chart)])
    tops = sorted(pt(p.top) for p in shapes(prs.slides[0], "Delta") + shapes(prs.slides[0], "Callout"))
    assert len(tops) == 3 and [p.text_frame.text for p in shapes(prs.slides[0], "Delta")] == ["2.0x", "+50%"]


def test_bar_chart_annotations_use_the_value_axis_across(tmp_path):
    chart = {"type": "bar", "categories": ["A", "B"], "series": [{"name": "s", "values": [40, 10]}],
             "annotations": [{"kind": "ref_line", "value": "mean", "label": "Average {value}"},
                             {"kind": "callout", "at": 0, "text": "Largest"}]}
    prs, _, _ = make(tmp_path, [chart_slide(chart)])
    slide = prs.slides[0]
    (frame,) = charts(slide)
    lo, hi = axis_scale(frame)
    x, y, w, h = plot_rect(frame)
    (ref,) = shapes(slide, "Reference line")
    assert pt(ref.left) == pytest.approx(x + w * (25 - lo) / (hi - lo), abs=0.05)
    assert shapes(slide, "Reference label")[0].text_frame.text == "Average 25"
    (pill,) = shapes(slide, "Callout")
    assert pt(pill.left) >= x + w and pt(pill.top) < y + h / 2  # beside the first (top) bar


def test_lint_annotation_rules(tmp_path):
    cats = [str(i) for i in range(10)]
    many = [{"kind": "callout", "at": c, "text": "x"} for c in cats[:4]]
    _, _, deck = make(tmp_path, [
        chart_slide({"categories": cats, "series": [{"name": "s", "values": list(range(10))}], "annotations": many}),
        chart_slide({"categories": ["a"], "series": [{"name": "s", "values": [1]}],
                     "annotations": [{"kind": "callout", "at": "zz"}, {"kind": "arrow"}]}),
        chart_slide({"type": "pie", "categories": ["a"], "series": [{"name": "s", "values": [1]}],
                     "annotations": [{"kind": "ref_line", "value": 1}]}),
        chart_slide({"categories": ["a", "b"], "number_format": "0%", "series": [{"name": "s", "values": [0.1, 0.2]}],
                     "annotations": [{"kind": "delta", "from": "a", "to": "b", "fmt": "pct"}]}),
    ])
    found = [(f["level"], f["slide"], f["message"]) for f in lint(deck)]
    assert any(lv == "WARN" and s == 1 and "4 callouts and deltas" in m for lv, s, m in found)
    assert any(lv == "WARN" and s == 1 and "10 categories" in m for lv, s, m in found)
    assert any(lv == "ERROR" and s == 2 and "'zz' is not a category" in m for lv, s, m in found)
    assert any(lv == "ERROR" and s == 2 and "unknown annotation kind 'arrow'" in m for lv, s, m in found)
    assert any(lv == "ERROR" and s == 3 and "not 'pie'" in m for lv, s, m in found)
    assert any(lv == "WARN" and s == 4 and "percentage points" in m for lv, s, m in found)


def test_number_and_delta_formatting():
    assert exhibits.format_number(0.456, "0%") == "46%"
    assert exhibits.format_number(1234.5, "#,##0.0") == "1,234.5"
    assert exhibits.format_number(1234567, '"$"#,##0,"k"') == "$1,235k"
    assert exhibits.format_number(-3, "0") == "−3"
    assert exhibits.format_number(2.5) == "2.5"
    assert exhibits.delta_text(0.42, 0.6, "pp", "0%") == "+18pp"
    assert exhibits.delta_text(42, 30, "pp", None, exhibits.WORDS["ko"]) == "−12%p"
    assert exhibits.delta_text(10, 25, "x", None) == "2.5x"
    assert exhibits.delta_text(10, 12, "pct", None) == "+20%"
    assert exhibits.delta_text(10, 12, "abs", "#,##0") == "+2"
    assert exhibits.axis_range(0, 58) == (0, 60, 20)  # about five gridline steps
    assert exhibits.axis_range(0, 0.64) == (0, 0.8, 0.2)


# -- 8. commentary ----------------------------------------------------------------------

def grid(start, span):
    gap = float(BRAND["layout.gutter"])
    cw = (BRAND.content_width - gap * (GRID - 1)) / GRID
    return BRAND.content_left + start * (cw + gap), span * cw + (span - 1) * gap


def test_takeaway_is_commentary_right_four_columns(tmp_path):
    ch = {"categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]}
    prs, _, _ = make(tmp_path, [chart_slide(ch, takeaway="The **reading**"),
                                chart_slide(ch, commentary={"text": "The **reading**"})])
    for slide in prs.slides:
        (frame,) = charts(slide)
        x, w = grid(0, 8)
        assert pt(frame.left) == pytest.approx(x, abs=0.01) and pt(frame.width) == pytest.approx(w, abs=0.01)
        (text,) = shapes(slide, "Commentary")
        assert pt(text.left) == pytest.approx(grid(8, 4)[0] + 12, abs=0.01)
    one, two = ([(s.name, s.left, s.top, s.width, s.height) for s in sl.shapes] for sl in prs.slides)
    assert one == two


@pytest.mark.parametrize("layout", ["chart", "table", "image", "diagram"])
@pytest.mark.parametrize("side,width", [("right", 3), ("left", 4)])
def test_commentary_column_on_exhibit_layouts(tmp_path, layout, side, width):
    from PIL import Image
    Image.new("RGB", (800, 450), "#888888").save(tmp_path / "shot.png")
    spec = {"layout": layout, "title": "A headline that states the result", "notes": "n",
            "commentary": {"bullets": ["One point", "Another point"], "side": side, "width": width}}
    spec.update({"chart": {"categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]},
                 "source": "Illustrative"} if layout == "chart" else
                {"columns": ["A", "B"], "rows": [["1", "2"]]} if layout == "table" else
                {"image": "shot.png", "alt": "A grey test image"} if layout == "image" else
                {"nodes": [{"id": "a", "label": "A", "col": 0, "row": 0}], "grid": {"cols": 1, "rows": 1}})
    prs, report, _ = make(tmp_path, [spec])
    slide = prs.slides[0]
    note_x, note_w = grid(0, width) if side == "left" else grid(GRID - width, width)
    ex_x, ex_w = grid(width, GRID - width) if side == "left" else grid(0, GRID - width)
    (rule,) = shapes(slide, "Commentary rule")
    assert pt(rule.left) == pytest.approx(note_x, abs=0.01)
    body = [s for s in slide.shapes if s.name in ("Chart", "Table", "Image", "Node a")]
    assert body
    for s in body:
        assert pt(s.left) >= ex_x - 0.01 and pt(s.left + s.width) <= ex_x + ex_w + 0.01
    assert not [w for w in report["warnings"] if w["level"] == "ERROR"]


def test_lint_commentary_rules(tmp_path):
    ch = {"categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]}
    _, _, deck = make(tmp_path, [chart_slide(ch, commentary={"text": "x", "width": 5, "side": "top"}, takeaway="y"),
                                 {"layout": "bullets", "title": "A headline that states the result",
                                  "bullets": ["a"], "commentary": {"text": "x"}, "notes": "n"}])
    found = [(f["level"], f["slide"], f["message"]) for f in lint(deck)]
    assert ("ERROR", 1, "commentary width is 3 or 4 grid columns") in found
    assert ("ERROR", 1, "commentary side must be right or left") in found
    assert any(lv == "WARN" and s == 1 and "takeaway is ignored" in m for lv, s, m in found)
    assert any(lv == "WARN" and s == 2 and "no effect on 'bullets'" in m for lv, s, m in found)


# -- 9. panels -------------------------------------------------------------------------

def panel(title, values, unit="Members", kind="column", **extra):
    return {"type": kind, "title": title, "unit": unit, "categories": ["x", "y", "z"],
            "series": [{"name": "s", "values": values}], **extra}


def test_panels_share_header_baseline_and_scale(tmp_path):
    callout = {"annotations": [{"kind": "callout", "at": "z", "text": "Record"}]}  # needs a taller top
    chart = {"panels": [panel("North", [41, 48, 56], **callout), panel("East", [3, 5, 8]),
                        panel("South", [29, 30, 33])]}
    prs, report, _ = make(tmp_path, [chart_slide(chart)])
    slide = prs.slides[0]
    frames = charts(slide)
    assert len(frames) == 3
    assert len({(f.top, f.height) for f in frames}) == 1
    assert len({tuple(manual_layout(f).values()) for f in frames}) == 1  # one plot rectangle
    assert {axis_scale(f) for f in frames} == {(0, 60)}  # same unit: one scale
    assert not any(f.chart.has_legend for f in frames)
    assert [s.text_frame.text for s in shapes(slide, "Panel title")] == ["North", "East", "South"]
    assert len(shapes(slide, "Panel unit")) == 3
    assert not report["warnings"]


def test_panels_with_different_units_keep_their_scales_and_one_key(tmp_path):
    two = [{"name": "Email", "values": [9, 8, 6]}, {"name": "Chat", "values": [2, 1.5, 1]}]
    chart = {"panels": [panel("Reply time", [0, 0, 0], unit="Hours", series=two, series_kind="nominal"),
                        panel("Reopened", [0, 0, 0], unit="Share", series=[
                            {"name": "Email", "values": [0.19, 0.17, 0.14]},
                            {"name": "Chat", "values": [0.12, 0.11, 0.1]}], series_kind="nominal",
                              number_format="0%")]}
    prs, _, deck = make(tmp_path, [chart_slide(chart)])
    slide = prs.slides[0]
    a, b = charts(slide)
    assert axis_scale(a) != axis_scale(b)
    assert manual_layout(a)["y"] == manual_layout(b)["y"] and manual_layout(a)["h"] == manual_layout(b)["h"]
    assert [s.text_frame.text for s in shapes(slide, "Key")] == ["Email", "Chat"]  # drawn once
    assert not any(f.chart.has_legend for f in (a, b))
    assert not [f for f in lint(deck) if "different colors" in f["message"]]


def test_lint_panel_rules(tmp_path):
    two = [{"name": "A", "values": [1, 2, 3]}, {"name": "B", "values": [3, 2, 1]}]
    chart = {"panels": [panel("One", [1, 2, 3], series=two), panel("Two", [1, 2, 3], series=two, series_kind="nominal"),
                        panel("3", [1, 2, 3]), panel("4", [1, 2, 3]), panel("5", [1, 2, 3])]}
    _, _, deck = make(tmp_path, [chart_slide(chart)])
    messages = [f["message"] for f in lint(deck)]
    assert any("5 chart panels" in m for m in messages)
    assert any("different colors" in m for m in messages)


# -- compatibility and dark slides ----------------------------------------------------------

def test_dark_slide_annotations_stay_readable(tmp_path):
    chart = {"type": "line", "categories": ["a", "b", "c"], "series": [{"name": "s", "values": [1, 2, 3]}],
             "annotations": [{"kind": "callout", "at": "b", "text": "Here"}]}
    prs, report, _ = make(tmp_path, [chart_slide(chart, dark=True)])
    (pill,) = shapes(prs.slides[0], "Callout")
    assert pill.fill.fore_color.rgb == BRAND.rgb("dark")
    (label,) = shapes(prs.slides[0], "End label")
    run = label.text_frame.paragraphs[0].runs[0]
    assert str(run.font.color.rgb) in (BRAND.hex("accent"), BRAND.hex("on_dark_body"))
    assert not [w for w in report["warnings"] if w["level"] == "ERROR"]


def test_example_deck_uses_every_exhibit_feature():
    spec = yaml.safe_load((SKILL / "examples/decks/exhibits.deck.yaml").read_text(encoding="utf-8"))
    text = yaml.safe_dump(spec)
    for key in ("survey", "question", "base", "as_of", "rounding", "highlight_series", "series_kind",
                "waterfall", "estimates", "ref_line", "band", "callout", "delta", "commentary",
                "takeaway", "panels"):
        assert key in text, key
    layouts = {s["layout"] for s in spec["slides"] if s.get("commentary")}
    assert layouts == {"chart", "table", "image", "diagram"}
    assert "illustrative" in text.lower()


def test_lint_cli_accepts_the_example():
    out = subprocess.run([sys.executable, str(SCRIPTS / "lint_deck.py"), str(SKILL / "examples/decks/exhibits.deck.yaml")],
                         capture_output=True, text=True)
    assert out.returncode == 0, out.stdout
    assert re.search(r"^0 error\(s\)", out.stdout, re.M)
