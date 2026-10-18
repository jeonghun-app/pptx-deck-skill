"""Regression tests: python -m pytest skills/pptx-deck/scripts/tests"""
import subprocess
import sys
from pathlib import Path

import pytest
import yaml
from pptx import Presentation
from pptx.util import Inches

SCRIPTS = Path(__file__).resolve().parents[1]
SKILL = SCRIPTS.parent
sys.path.insert(0, str(SCRIPTS))

from deckkit import build, load_brand, contrast  # noqa: E402
from deckkit import measure  # noqa: E402
from deckkit.layouts import LAYOUTS  # noqa: E402

SAMPLES = sorted((SKILL / "examples" / "decks").glob("*.deck.yaml"))


def run(*args):
    return subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True)


@pytest.mark.parametrize("deck", SAMPLES, ids=lambda p: p.stem)
def test_samples_lint_and_build_clean(deck, tmp_path):
    lint = run(SCRIPTS / "lint_deck.py", deck)
    assert lint.returncode == 0, lint.stdout
    report = build(deck, tmp_path / "deck.pptx")
    assert [w for w in report["warnings"] if w["level"] == "ERROR"] == []
    prs = Presentation(str(tmp_path / "deck.pptx"))
    spec = yaml.safe_load(deck.read_text(encoding="utf-8"))
    assert len(prs.slides) == len(spec["slides"])
    props = prs.core_properties
    assert props.last_modified_by == props.author == spec["meta"].get("author", "")
    for slide, s in zip(prs.slides, spec["slides"]):
        assert slide.notes_slide.notes_text_frame.text.strip() == str(s["notes"]).strip()


def test_sample_covers_every_layout():
    spec = yaml.safe_load((SKILL / "examples/decks/sample.deck.yaml").read_text(encoding="utf-8"))
    used = {s["layout"] for s in spec["slides"]}
    ko = yaml.safe_load((SKILL / "examples/decks/ko-sample.deck.yaml").read_text(encoding="utf-8"))
    used |= {s["layout"] for s in ko["slides"]}
    assert set(LAYOUTS) == used


def test_partial_brand_inherits_defaults():
    brand = load_brand(SKILL / "examples/brands/warm-editorial/brand.yaml")
    assert brand.hex("accent") == "9C4422"
    assert brand["type.lead"] == load_brand()["type.lead"]  # inherited
    assert brand.get("layout.rule_y") is None


BRANDS = [None] + sorted((SKILL / "examples" / "brands").glob("*/brand.yaml"))


@pytest.mark.parametrize("path", BRANDS, ids=lambda p: "default" if p is None else p.parent.name)
def test_shipped_brands_pass_text_contrast(path):
    b = load_brand(path)
    for bg in ("surface", "surface_alt"):
        for token in ("ink", "body", "muted", "accent"):  # accent is used for eyebrow text
            assert contrast(b.hex(token), b.hex(bg)) >= 4.5, (token, bg)
    for role in ("cover", "section", "closing"):
        bg = b.get(f"backgrounds.{role}.color")
        if bg and b.is_dark(bg):
            assert contrast(b.hex("on_dark_body"), b.hex(bg)) >= 4.5, role


def test_tokens_drive_output(tmp_path):
    brand = tmp_path / "brand.yaml"
    brand.write_text("colors: {accent: 123456}\n", encoding="utf-8")
    deck = tmp_path / "d.deck.yaml"
    deck.write_text(yaml.safe_dump({
        "meta": {"audience": "a", "objective": "o", "desired_action": "d"},
        "slides": [{"layout": "bullets", "eyebrow": "x", "title": "A headline that states a result",
                    "bullets": ["one"], "notes": "n"}]}), encoding="utf-8")
    build(deck, tmp_path / "o.pptx", brand)
    xml = Presentation(str(tmp_path / "o.pptx")).slides[0].shapes._spTree.xml
    assert "123456" in xml


def test_overflow_is_reported(tmp_path):
    deck = tmp_path / "d.deck.yaml"
    deck.write_text(yaml.safe_dump({
        "meta": {"audience": "a", "objective": "o", "desired_action": "d"},
        "slides": [{"layout": "bullets", "title": "Too much text for one slide to hold",
                    "bullets": ["a long sentence that keeps going " * 6] * 8, "notes": "n"}]}),
        encoding="utf-8")
    report = build(deck, tmp_path / "o.pptx")
    assert any(w["level"] == "ERROR" and "below the body limit" in w["message"] for w in report["warnings"])


def test_lint_blocks_missing_contract_and_emoji(tmp_path):
    deck = tmp_path / "d.deck.yaml"
    deck.write_text(yaml.safe_dump({"meta": {}, "slides": [
        {"layout": "bullets", "title": "Launch 🚀 went well for every team", "bullets": ["x"]}]},
        allow_unicode=True), encoding="utf-8")
    out = run(SCRIPTS / "lint_deck.py", deck)
    assert out.returncode == 1
    assert "meta.objective is empty" in out.stdout and "emoji" in out.stdout


def test_template_mode_round_trip(tmp_path):
    tpl = tmp_path / "company.pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.slides.add_slide(prs.slide_layouts[0])  # existing slides must be dropped
    prs.save(str(tpl))
    brand_file = tmp_path / "brand" / "brand.yaml"
    out = run(SCRIPTS / "extract_brand.py", tpl, "-o", brand_file)
    assert out.returncode == 0, out.stderr
    brand = load_brand(brand_file)
    assert brand.get("template.layouts.content") == "Title Only"
    assert (brand_file.parent / "assets" / "company.pptx").exists()
    report = build(SKILL / "examples/decks/sample.deck.yaml", tmp_path / "t.pptx", brand_file)
    assert [w for w in report["warnings"] if w["level"] == "ERROR"] == []
    built = Presentation(str(tmp_path / "t.pptx"))
    assert built.slides[0].slide_layout.name == "Title Slide"
    assert built.slides[0].shapes.title.text.startswith("Catch release blockers")


def test_measure_wraps_long_text():
    assert measure.count_lines("short", 16, 400) == 1
    assert measure.count_lines("word " * 60, 16, 400) > 3
    assert measure.count_lines("가나다라마바사아자차" * 5, 16, 200) > 1


def test_image_layout_and_missing_asset(tmp_path):
    from PIL import Image
    Image.new("RGB", (800, 450), "#888888").save(tmp_path / "shot.png")
    deck = tmp_path / "d.deck.yaml"
    deck.write_text(yaml.safe_dump({
        "meta": {"audience": "a", "objective": "o", "desired_action": "d"},
        "slides": [
            {"layout": "image", "title": "The new screen removes two steps", "image": "shot.png",
             "alt": "Screen", "caption": "c", "notes": "n"},
            {"layout": "image", "title": "This image file does not exist yet", "image": "nope.png",
             "alt": "x", "notes": "n"}]}), encoding="utf-8")
    report = build(deck, tmp_path / "o.pptx")
    errors = [w for w in report["warnings"] if w["level"] == "ERROR"]
    assert len(errors) == 1 and errors[0]["slide"] == 2


def test_cli_lints_before_building_and_writes_clean_app_props(tmp_path):
    bad = tmp_path / "bad.deck.yaml"
    bad.write_text("meta: {}\nslides:\n- {layout: bullets, title: x, bullets: [a]}\n", encoding="utf-8")
    out = run(SCRIPTS / "build_deck.py", bad, "-o", tmp_path / "bad.pptx")
    assert out.returncode == 1 and not (tmp_path / "bad.pptx").exists()
    good = run(SCRIPTS / "build_deck.py", SKILL / "examples/decks/ko-sample.deck.yaml", "-o", tmp_path / "ko.pptx")
    assert good.returncode == 0, good.stdout
    import zipfile
    app = zipfile.ZipFile(tmp_path / "ko.pptx").read("docProps/app.xml").decode()
    assert "<Slides>6</Slides>" in app and "16:9" in app
