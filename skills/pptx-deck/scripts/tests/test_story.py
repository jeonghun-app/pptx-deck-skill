"""Reading structure: measure, reading mode, summary, tracker, bands, Korean notation.

python -m pytest skills/pptx-deck/scripts/tests/test_story.py
"""
import sys
from pathlib import Path

import pytest
import yaml
from pptx import Presentation

SCRIPTS = Path(__file__).resolve().parents[1]
SKILL = SCRIPTS.parent
sys.path.insert(0, str(SCRIPTS))

from deckkit import build, load_brand, story  # noqa: E402
from deckkit import measure  # noqa: E402
from lint_deck import lint  # noqa: E402

META = {"audience": "a", "objective": "o", "desired_action": "d"}
PT = 12700  # EMU per point


def write_deck(tmp_path, slides, name="d", **meta):
    deck = tmp_path / f"{name}.deck.yaml"
    for s in slides:
        s.setdefault("notes", "n")
    deck.write_text(yaml.safe_dump({"meta": {**META, **meta}, "slides": slides}, allow_unicode=True),
                    encoding="utf-8")
    return deck


def build_deck(tmp_path, slides, brand=None, **meta):
    deck = write_deck(tmp_path, slides, **meta)
    report = build(deck, tmp_path / "o.pptx", brand)
    return Presentation(str(tmp_path / "o.pptx")), report


def shapes(slide, prefix):
    return [s for s in slide.shapes if s.name.startswith(prefix)]


def top(shape):
    return shape.top / PT


def width(shape):
    return shape.width / PT


def messages(findings, level=None):
    return [f["message"] for f in findings if level is None or f["level"] == level]


def lint_slides(tmp_path, slides, **meta):
    return lint(write_deck(tmp_path, slides, **meta))


SECTIONS = [
    {"layout": "section", "number": 1, "title": "Where time goes", "short": "Diagnosis",
     "subtitle": "Waiting, not work, takes most of the time."},
    {"layout": "bullets", "eyebrow": "auto", "title": "Three handoffs hold most of the waiting time",
     "bullets": ["one", "two"]},
    {"layout": "section", "number": 2, "title": "What we change", "short": "Changes"},
    {"layout": "bullets", "eyebrow": "auto", "title": "One owner per account removes the handoffs",
     "bullets": ["one", "two"]},
]


# -- 1. reading measure -----------------------------------------------------------

def test_chars_per_line_counts_cjk_as_about_two():
    assert measure.chars_per_line(16, measure.measure_width(16, 75)) == pytest.approx(75)
    width_ = measure.measure_width(16, 75)
    latin = "The pilot continues only if the measures improve without slowing merges."
    assert len(latin) > 70 and measure.count_lines(latin, 16, width_) == 1
    assert measure.count_lines("가" * 36, 16, width_) == 1
    assert measure.count_lines("가" * 75, 16, width_) >= 2  # a Hangul glyph counts about two


def test_lead_and_bullets_keep_the_reading_measure(tmp_path):
    long = "A reading sentence that is long enough to wrap across more than one line " * 2
    prs, _ = build_deck(tmp_path, [{"layout": "bullets", "title": "A headline that states a result",
                                    "lead": long, "bullets": [long, "short"]}])
    slide = prs.slides[0]
    b = load_brand()
    cap = measure.measure_width(b["type.lead"], b["type.max_measure"])
    lead = shapes(slide, "Lead")[0]
    assert width(lead) == pytest.approx(cap, abs=0.5) and cap < b.content_width
    bullets = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.startswith("A reading")
               and s.name != "Lead"]
    assert bullets and all(width(s) <= cap + 0.5 for s in bullets)
    assert all(lead.left == s.left for s in bullets)  # aligned left, not centered


def test_measure_cap_can_be_switched_off(tmp_path):
    brand = tmp_path / "brand.yaml"
    brand.write_text("type: {max_measure: null}\n", encoding="utf-8")
    prs, _ = build_deck(tmp_path, [{"layout": "bullets", "title": "A headline that states a result",
                                    "lead": "Lead text", "bullets": ["one"]}], brand=brand)
    assert width(shapes(prs.slides[0], "Lead")[0]) == pytest.approx(load_brand().content_width)


def test_narrow_columns_are_not_affected_by_the_measure(tmp_path):
    prs, _ = build_deck(tmp_path, [{"layout": "two_column", "title": "A headline that states a result",
                                    "left": {"heading": "L", "bullets": ["one"]},
                                    "right": {"heading": "R", "bullets": ["two"]}}])
    col_w = (load_brand().content_width - 36) / 2
    one = next(s for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text == "one")
    assert width(one) == pytest.approx(col_w, abs=0.5)


def test_lint_warns_on_over_wide_multi_line_caption(tmp_path):
    from PIL import Image
    Image.new("RGB", (80, 45), "#888888").save(tmp_path / "shot.png")
    caption = "A caption that keeps going well past the reading measure of the slide body text " * 2
    found = lint_slides(tmp_path, [{"layout": "image", "title": "The screen removes two of the steps",
                                    "image": "shot.png", "alt": "x", "caption": caption}])
    assert any("caption runs over several lines" in m for m in messages(found, "WARN"))
    found = lint_slides(tmp_path, [{"layout": "image", "title": "The screen removes two of the steps",
                                    "image": "shot.png", "alt": "x", "caption": "One short line"}])
    assert not any("caption runs" in m for m in messages(found))


# -- 2. reading mode --------------------------------------------------------------

def body_tops(prs):
    return [top(shapes(s, "Lead")[0]) for s in prs.slides]


def test_read_mode_reserves_two_title_lines_on_every_content_slide(tmp_path):
    slides = [{"layout": "bullets", "title": "Short headline that states a result", "lead": "L",
               "bullets": ["a"]},
              {"layout": "chart", "title": "A much longer headline that needs two lines in the band "
                                           "to carry both the claim and its consequence",
               "lead": "L", "chart": {"categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]},
               "source": "s"}]
    present, _ = build_deck(tmp_path, [dict(s) for s in slides])
    read, report = build_deck(tmp_path, [dict(s) for s in slides], mode="read")
    b = load_brand()
    line = b["type.title"] * 1.2 * b["type.heading_line_spacing"]
    assert body_tops(present) == pytest.approx([b["layout.body_top"]] * 2)
    assert body_tops(read) == pytest.approx([b["layout.body_top"] + line] * 2)  # fixed band
    rule = [top(shapes(s, "Rule")[0]) for s in read.slides]
    assert rule == pytest.approx([b["layout.rule_y"] + line] * 2)
    assert not [w for w in report["warnings"] if "headline wraps" in w["message"]]


def test_brand_that_already_reserves_two_lines_is_not_shifted_twice():
    brand = load_brand()
    two = story.with_layout(brand, title_lines=2)
    assert story.frame_brand(two, {"mode": "read"}) is two
    assert story.frame_brand(brand, {}) is brand
    assert story.title_lines(brand, {"mode": "read"}) == 2


def test_lint_headline_ceiling_follows_the_mode(tmp_path):
    two = "Onboarding is slow because accounts wait between teams so one owner should come first"
    slides = [{"layout": "bullets", "title": two, "bullets": ["a"]}]
    assert any("past one line" in m for m in messages(lint_slides(tmp_path, slides)))
    read = messages(lint_slides(tmp_path, slides, mode="read"))
    assert not any("wraps past" in m for m in read)
    long = " ".join(["word"] * 30)
    read = messages(lint_slides(tmp_path, [{"layout": "bullets", "title": long, "bullets": ["a"]}],
                                mode="read"))
    assert any("30 words" in m for m in read)
    bad = lint_slides(tmp_path, slides, mode="slides")
    assert any("meta.mode" in m for m in messages(bad, "ERROR"))


# -- 3. summary ---------------------------------------------------------------------

def summary_deck(**extra):
    return [{"layout": "summary", "title": "Waiting between teams, not work, makes onboarding slow",
             "blocks": [{"section": 1, "points": ["Eleven days, two of them work"]},
                        {"section": 2, "points": ["One owner per account"]}],
             "so_what": "Approve a six-week trial", **extra}] + [dict(s) for s in SECTIONS]


def text_of(slide):
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame]


def test_summary_pulls_section_titles_and_theses(tmp_path):
    prs, report = build_deck(tmp_path, summary_deck())
    texts = text_of(prs.slides[0])
    assert "Where time goes" in texts and "What we change" in texts
    assert "Waiting, not work, takes most of the time." in texts  # section 1 subtitle as thesis
    assert "01" in texts and "02" in texts
    band = shapes(prs.slides[0], "So what")[0]
    assert band.fill.fore_color.rgb == load_brand().rgb("surface_alt")
    assert [w for w in report["warnings"] if w["level"] == "ERROR"] == []


def test_summary_band_sits_below_the_blocks(tmp_path):
    prs, _ = build_deck(tmp_path, summary_deck())
    slide = prs.slides[0]
    band = shapes(slide, "So what")[0]
    point = next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text == "One owner per account")
    assert top(point) + point.height / PT < top(band)
    assert top(band) + band.height / PT <= load_brand()["layout.body_bottom"] + 0.5


def test_summary_reference_out_of_range_is_an_error(tmp_path):
    slides = summary_deck()
    slides[0]["blocks"][1]["section"] = 5
    _, report = build_deck(tmp_path, slides)
    assert any(w["level"] == "ERROR" and "section 5" in w["message"] for w in report["warnings"])
    assert any("section 5" in m for m in messages(lint_slides(tmp_path, slides), "ERROR"))


def test_lint_summary_heading_must_match_its_section(tmp_path):
    slides = summary_deck()
    slides[0]["blocks"][0]["heading"] = "Something else entirely"
    assert any("differs from section 1" in m for m in messages(lint_slides(tmp_path, slides), "WARN"))
    slides[0]["blocks"][0]["heading"] = "where time goes"  # same wording, different case
    assert not any("differs" in m for m in messages(lint_slides(tmp_path, slides)))


def test_lint_long_deck_without_summary(tmp_path):
    slides = [{"layout": "bullets", "title": f"Finding number {k} changes the plan", "bullets": ["a"]}
              for k in range(10)]
    assert any("no summary" in m for m in messages(lint_slides(tmp_path, slides), "WARN"))
    slides[0] = summary_deck()[0]
    slides[0]["blocks"] = [{"heading": "One", "points": ["a"]}, {"heading": "Two", "points": ["b"]}]
    assert not any("no summary" in m for m in messages(lint_slides(tmp_path, slides)))


# -- 4. tracker and eyebrow: auto ----------------------------------------------------

def deck_with_sections():
    return ([{"layout": "cover", "title": "Cover"},
             {"layout": "bullets", "eyebrow": "Before", "title": "A slide before the first section",
              "bullets": ["a"]}]
            + [dict(s) for s in SECTIONS] + [{"layout": "closing"}])


def test_tracker_tabs_replace_the_eyebrow_inside_sections(tmp_path):
    prs, report = build_deck(tmp_path, deck_with_sections(), tracker="tabs")
    slides = list(prs.slides)
    assert shapes(slides[1], "Tracker") == [] and shapes(slides[1], "Eyebrow")  # before section 1
    for k, slide in ((0, slides[3]), (1, slides[5])):
        tabs = shapes(slide, "Tracker tab")
        assert [t.text_frame.text for t in tabs] == ["01 DIAGNOSIS", "02 CHANGES"]
        accent = load_brand().rgb("accent")
        colors = [t.text_frame.paragraphs[0].runs[0].font.color.rgb for t in tabs]
        assert [c == accent for c in colors] == [i == k for i in range(2)]
        assert shapes(slide, "Eyebrow") == []
        assert top(shapes(slide, "Tracker rule")[0]) < load_brand()["layout.title_y"]
    for structural in (slides[0], slides[2], slides[4], slides[6]):
        assert shapes(structural, "Tracker") == []
    assert [w for w in report["warnings"] if "tracker" in w["message"]] == []


def test_tracker_dots_keep_an_auto_eyebrow(tmp_path):
    prs, _ = build_deck(tmp_path, deck_with_sections(), tracker="dots")
    slide = prs.slides[5]
    dots = shapes(slide, "Tracker dot")
    assert len(dots) == 2
    assert dots[1].fill.fore_color.rgb == load_brand().rgb("accent")
    eyebrow = shapes(slide, "Eyebrow")[0]
    assert eyebrow.text_frame.text == "02 · CHANGES"
    assert eyebrow.left + eyebrow.width <= dots[0].left  # never under the dots


def test_tracker_dots_stay_clear_of_a_top_right_logo(tmp_path):
    from PIL import Image
    Image.new("RGB", (128, 40), "#333333").save(tmp_path / "logo.png")
    brand = tmp_path / "brand.yaml"
    brand.write_text("chrome: {logo: {path: logo.png, position: top-right, width: 64}}\n", encoding="utf-8")
    prs, _ = build_deck(tmp_path, deck_with_sections(), brand=brand, tracker="dots")
    dots = shapes(prs.slides[3], "Tracker dot")
    right = max(d.left + d.width for d in dots) / PT
    assert right <= load_brand().content_right - 64 - 16 + 0.5


def test_tracker_from_brand_and_deck_override(tmp_path):
    brand = tmp_path / "brand.yaml"
    brand.write_text("chrome: {tracker: dots}\n", encoding="utf-8")
    prs, _ = build_deck(tmp_path, deck_with_sections(), brand=brand)
    assert shapes(prs.slides[3], "Tracker dot")
    prs, _ = build_deck(tmp_path, deck_with_sections(), brand=brand, tracker="none")
    assert shapes(prs.slides[3], "Tracker") == []


def test_tracker_falls_back_and_warns(tmp_path):
    many = []
    for k in range(7):
        many += [{"layout": "section", "number": k + 1, "title": f"Part {k + 1}"},
                 {"layout": "bullets", "title": f"Finding {k + 1} changes the plan", "bullets": ["a"]}]
    prs, report = build_deck(tmp_path, many, tracker="tabs")
    assert any("drawn as dots" in w["message"] for w in report["warnings"])
    assert len(shapes(prs.slides[1], "Tracker dot")) == 7
    _, report = build_deck(tmp_path, [dict(SECTIONS[0]), dict(SECTIONS[1])], tracker="tabs")
    assert any("two or more section" in w["message"] for w in report["warnings"])
    assert any("too many tabs" in m for m in messages(lint_slides(tmp_path, many, tracker="tabs")))
    assert any("tracker 'x'" in m for m in messages(lint_slides(tmp_path, many, tracker="x"), "ERROR"))


def test_tracker_warns_once_when_a_tab_label_is_too_long(tmp_path):
    slides = deck_with_sections()
    slides[2]["short"] = "A section name that is far too long to be read in one tab of the tracker band"
    _, report = build_deck(tmp_path, slides, tracker="tabs")
    hits = [w for w in report["warnings"] if "wider than its tab" in w["message"]]
    assert len(hits) == 1 and hits[0]["slide"] == 3


def test_eyebrow_auto_before_any_section(tmp_path):
    slides = [{"layout": "bullets", "eyebrow": "auto", "title": "A slide outside every section",
               "bullets": ["a"]}]
    prs, report = build_deck(tmp_path, [dict(s) for s in slides])
    assert shapes(prs.slides[0], "Eyebrow") == []
    assert any("eyebrow: auto" in w["message"] for w in report["warnings"])
    assert any("eyebrow: auto" in m for m in messages(lint_slides(tmp_path, slides), "WARN"))


# -- 5. lead band ---------------------------------------------------------------------

def test_lead_band_under_a_question_headline(tmp_path):
    prs, _ = build_deck(tmp_path, [{"layout": "bullets", "title": "Which step holds an account longest?",
                                    "lead": "The wait before setup.", "lead_style": "band",
                                    "bullets": ["a"]}])
    slide = prs.slides[0]
    b = load_brand()
    band = shapes(slide, "Lead band")[0]
    assert band.fill.fore_color.rgb == b.rgb("surface_alt")
    assert width(band) == pytest.approx(b.content_width)
    assert top(band) == pytest.approx(b["layout.body_top"])
    rule = shapes(slide, "Lead band rule")[0]
    assert rule.fill.fore_color.rgb == b.rgb("accent") and rule.left == band.left
    first = next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text == "a")
    assert top(first) >= top(band) + band.height / PT


def test_lint_lead_style(tmp_path):
    base = {"layout": "bullets", "title": "Which step holds an account the longest?", "bullets": ["a"]}
    assert any("lead_style 'box'" in m for m in messages(
        lint_slides(tmp_path, [{**base, "lead": "x", "lead_style": "box"}]), "ERROR"))
    assert any("without a lead" in m for m in messages(
        lint_slides(tmp_path, [{**base, "lead_style": "band"}]), "WARN"))


# -- 6. next steps band ------------------------------------------------------------------

def test_next_steps_reserves_its_height_before_the_body(tmp_path):
    slide_spec = {"layout": "chart", "title": "Setup waits longest of all the steps",
                  "chart": {"categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]},
                  "source": "Illustrative", "next_steps": {"items": ["Name owners", "Replace the inbox"]}}
    prs, report = build_deck(tmp_path, [slide_spec])
    slide = prs.slides[0]
    b = load_brand()
    band = shapes(slide, "Next steps")[0]
    assert top(band) + band.height / PT == pytest.approx(b["layout.body_bottom"])
    source = shapes(slide, "Source")[0]
    chart = shapes(slide, "Chart")[0]
    assert top(source) + source.height / PT <= top(band) - b["layout.gutter"] + 0.5
    assert top(chart) + chart.height / PT <= top(source) + 0.5
    texts = text_of(slide)
    assert "What this means" in [t.title() for t in texts] or "WHAT THIS MEANS" in texts
    assert "01" in texts and "02" in texts
    assert [w for w in report["warnings"] if w["level"] == "ERROR"] == []


def test_next_steps_shrinks_the_overflow_limit(tmp_path):
    bullets = ["A bullet that takes some room on the slide body and wraps to a second line "
               "at the reading measure"] * 5
    plain = {"layout": "bullets", "title": "A headline that states a result", "lead": "Lead",
             "bullets": bullets}
    _, report = build_deck(tmp_path, [dict(plain)])
    assert not [w for w in report["warnings"] if w["level"] == "ERROR"]
    crowded = {**plain, "next_steps": {"items": ["One step to take next " * 3] * 4}}
    _, report = build_deck(tmp_path, [crowded])
    assert any(w["level"] == "ERROR" and "body limit" in w["message"] for w in report["warnings"])


def test_next_steps_label_default_follows_language(tmp_path):
    spec = {"layout": "bullets", "title": "대기 시간의 절반은 첫 단계에서 발생", "bullets": ["a"],
            "next_steps": ["담당자 지정", "점검표 도입"]}
    prs, _ = build_deck(tmp_path, [spec], language="ko")
    assert "시사점" in text_of(prs.slides[0])


def test_lint_next_steps_frequency_and_count(tmp_path):
    slides = [{"layout": "bullets", "title": f"Finding number {k} changes the plan", "bullets": ["a"],
               **({"next_steps": ["one", "two"]} if k < 2 else {})} for k in range(4)]
    assert any("next_steps on 2 of 4" in m for m in messages(lint_slides(tmp_path, slides), "WARN"))
    slides[1].pop("next_steps")
    assert not any("next_steps on" in m for m in messages(lint_slides(tmp_path, slides)))
    slides[0]["next_steps"] = ["only one"]
    assert any("has 1 items" in m for m in messages(lint_slides(tmp_path, slides), "WARN"))


# -- 7. Korean conventions ------------------------------------------------------------------

def test_lint_korean_headline_endings(tmp_path):
    slides = [{"layout": "bullets", "title": "검토 요청의 절반이 하루 이상 대기합니다", "bullets": ["a"]},
              {"layout": "bullets", "title": "검토 요청의 절반이 하루 이상 대기", "bullets": ["a"]},
              {"layout": "bullets", "title": "검토 대기가 리드타임을 결정한다", "bullets": ["a"]},
              {"layout": "bullets", "title": "대기의 대부분은 실제 작업이 아니다", "bullets": ["a"]},
              {"layout": "bullets", "title": "검토 요청의 절반은 하루 이상 대기 중입니다", "bullets": ["a"]}]
    found = lint_slides(tmp_path, slides, language="ko")
    info = [f for f in found if f["level"] == "INFO"]
    assert [f["slide"] for f in info] == [1, 5]
    assert not any(f["level"] == "ERROR" for f in found)


def test_lint_korean_percentage_points_and_source_prefixes(tmp_path):
    slides = [{"layout": "stats", "title": "승인율이 전년 대비 3%p 상승", "source": "자료: 내부 집계",
               "stats": [{"value": "3%p", "label": "승인율 변화"}]},
              {"layout": "stats", "title": "반려율은 2%포인트 하락", "source": "출처: 내부 집계",
               "stats": [{"value": "2%포인트", "label": "반려율 변화"}]}]
    warns = messages(lint_slides(tmp_path, slides, language="ko"), "WARN")
    assert any("%p" in m and "%포인트" in m for m in warns)
    assert any("more than one prefix" in m for m in warns)
    slides[1]["title"], slides[1]["source"] = "반려율은 2%p 하락", "자료: 내부 집계"
    slides[1]["stats"] = [{"value": "2%p", "label": "반려율 변화"}]
    warns = messages(lint_slides(tmp_path, slides, language="ko"), "WARN")
    assert not any("%포인트" in m or "prefix" in m for m in warns)


def test_unprefixed_sources_count_as_the_default_label(tmp_path):
    slides = [{"layout": "stats", "title": "승인율이 전년 대비 3%p 상승", "source": "내부 집계",
               "stats": [{"value": "3%p", "label": "승인율 변화"}]},
              {"layout": "stats", "title": "반려율은 2%p 하락 추세", "source": {"text": "자료: 내부 집계"},
               "stats": [{"value": "2%p", "label": "반려율 변화"}]}]
    warns = messages(lint_slides(tmp_path, slides, language="ko"), "WARN")
    assert any("출처" in m and "자료" in m for m in warns)  # default 출처 mixed with 자료


def test_korean_sample_follows_the_headline_default():
    found = lint(SKILL / "examples/decks/ko-sample.deck.yaml")
    assert [f for f in found if f["level"] == "INFO"] == []


def test_story_example_uses_every_story_feature():
    spec = yaml.safe_load((SKILL / "examples/decks/story.deck.yaml").read_text(encoding="utf-8"))
    meta, slides = spec["meta"], spec["slides"]
    assert meta["mode"] == "read" and meta["tracker"] in ("tabs", "dots")
    keys = {k for s in slides for k in s}
    assert {"next_steps", "lead_style", "short"} <= keys
    assert any(s.get("eyebrow") == "auto" for s in slides)
    assert any(s["layout"] == "summary" and s.get("so_what") for s in slides)
