"""Framework layouts: matrix, shift, symbol tables, chevron process, open cards, stat grids.

    python -m pytest skills/pptx-deck/scripts/tests/test_frameworks.py
"""
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SCRIPTS = Path(__file__).resolve().parents[1]
SKILL = SCRIPTS.parent
sys.path.insert(0, str(SCRIPTS))

from deckkit import build, load_brand  # noqa: E402
from deckkit.frameworks import normalize_symbol  # noqa: E402
from lint_deck import lint  # noqa: E402

META = {"audience": "a", "objective": "o", "desired_action": "d"}
PT = 12700  # EMU per point


def deck(tmp_path, *slides, **meta):
    path = tmp_path / "d.deck.yaml"
    slides = [{"notes": "n", "title": "A headline that states the result", **s} for s in slides]
    path.write_text(yaml.safe_dump({"meta": {**META, **meta}, "slides": slides}, allow_unicode=True),
                    encoding="utf-8")
    return path


def built(tmp_path, *slides, brand=None, **meta):
    path = deck(tmp_path, *slides, **meta)
    report = build(path, tmp_path / "o.pptx", brand)
    return report, Presentation(str(tmp_path / "o.pptx"))


def shapes(prs, n=0):
    return {s.name: s for s in prs.slides[n].shapes}


def texts(prs, n=0):
    out = []
    for s in prs.slides[n].shapes:
        if s.has_text_frame:
            out.append(s.text_frame.text)
        if s.has_table:
            out += [c.text_frame.text for row in s.table.rows for c in row.cells]
    return out


def findings(tmp_path, *slides, level=None, **meta):
    out = lint(deck(tmp_path, *slides, **meta))
    return [f["message"] for f in out if level is None or f["level"] == level]


def errors(report):
    return [w["message"] for w in report["warnings"] if w["level"] == "ERROR"]


def centre(shape):
    return (shape.left + shape.width / 2) / PT, (shape.top + shape.height / 2) / PT


QUADRANTS = [{"name": "Plan"}, {"name": "Start here", "highlight": True}, {"name": "Avoid"},
             {"name": "Quick fixes"}]
AXES = {"x": {"label": "Ease", "low": "Hard", "high": "Easy"},
        "y": {"label": "Value", "low": "Low", "high": "High"}}


# -- matrix ------------------------------------------------------------------

def test_matrix_qualitative_lists_items_in_named_cells_and_tags_the_basis(tmp_path):
    report, prs = built(tmp_path, {"layout": "matrix", **AXES, "quadrants": QUADRANTS, "items": [
        {"id": "A", "label": "Contract review", "cell": 0},
        {"id": "B", "label": "Invoice matching", "cell": 1}]})
    assert errors(report) == []
    sh = shapes(prs)
    assert [sh[f"Cell {i} name"].text_frame.text for i in range(1, 5)] == [
        "Plan", "Start here", "Avoid", "Quick fixes"]
    brand = load_brand()
    assert str(sh["Cell 2"].fill.fore_color.rgb) == brand.tint("accent", 0.9)  # the one highlight
    assert str(sh["Cell 1"].fill.fore_color.rgb) == brand.hex("surface_alt")
    assert sh["Tag"].text_frame.text == "Qualitative placement"
    assert "B  Invoice matching" in texts(prs)
    assert not any(name.startswith("Dot") for name in sh)  # no fake precision
    # the highlighted cell sits top-right: right of and level with cell 1
    assert sh["Cell 2"].left > sh["Cell 1"].left and sh["Cell 2"].top == sh["Cell 1"].top
    assert "Hard" in texts(prs) and "Easy" in texts(prs) and sh["Y axis center"].rotation == 270


def test_matrix_measured_places_dots_from_values_and_draws_thresholds(tmp_path):
    report, prs = built(tmp_path, {
        "layout": "matrix", "basis": "measured", "source": "Scores, illustrative",
        "x": {**AXES["x"], "min": 0, "max": 100}, "y": {**AXES["y"], "min": 0, "max": 20},
        "thresholds": {"x": 60, "y": 8}, "quadrants": QUADRANTS,
        "items": [{"id": "L", "label": "Left", "x": 0, "y": 10}, {"id": "R", "label": "Right", "x": 100, "y": 10},
                  {"id": "M", "label": "Mid", "x": 50, "y": 20}]})
    assert errors(report) == []
    sh = shapes(prs)
    left, right, mid = (centre(sh[f"Dot {k}"]) for k in "LRM")
    assert mid[0] == pytest.approx((left[0] + right[0]) / 2, abs=0.5)
    gx0 = sh["Cell 1"].left / PT
    gx1 = (sh["Cell 2"].left + sh["Cell 2"].width) / PT
    assert left[0] == pytest.approx(gx0, abs=0.5) and right[0] == pytest.approx(gx1, abs=0.5)
    assert mid[1] == pytest.approx(sh["Cell 1"].top / PT, abs=0.5)  # y = max sits on the top edge
    line_x = sh["Threshold x"].left / PT
    assert line_x == pytest.approx(gx0 + 0.6 * (gx1 - gx0), abs=0.5)  # cells split at the threshold
    assert sh["Cell 2"].left / PT == pytest.approx(line_x, abs=0.5)
    ln = sh["Threshold x"].line._get_or_add_ln()
    assert ln.find(qn("a:prstDash")).get("val") == "dash"
    assert "Tag" not in sh and "60" in texts(prs) and "Low · 0" in texts(prs)


def test_matrix_commentary_takes_a_side_column(tmp_path):
    _, prs = built(tmp_path, {"layout": "matrix", **AXES, "quadrants": QUADRANTS, "items": [],
                              "commentary": {"side": "left", "width": 3, "bullets": ["One", "Two"]}})
    sh = shapes(prs)
    brand = load_brand()
    assert sh["Cell 1"].left / PT > brand.content_left + brand.content_width / 4
    _, prs = built(tmp_path, {"layout": "matrix", **AXES, "quadrants": QUADRANTS, "items": [],
                              "takeaway": "Two of six are ready."})
    sh = shapes(prs)
    assert "Two of six are ready." in texts(prs)
    assert (sh["Cell 2"].left + sh["Cell 2"].width) / PT < brand.content_right - brand.content_width / 4


def test_matrix_lint(tmp_path):
    base = {"layout": "matrix", **AXES, "quadrants": QUADRANTS}
    errs = findings(tmp_path, {**base, "basis": "measured", "items": [{"label": "a", "x": 1, "y": 1}]},
                    level="ERROR")
    assert any("without a source" in m for m in errs)
    warns = findings(tmp_path, {**base, "items": [{"label": "a", "cell": 0, "x": 0.2, "y": 0.4}]},
                     level="WARN")
    assert any("basis is qualitative" in m for m in warns)
    two = [dict(q, highlight=True) for q in QUADRANTS[:2]] + QUADRANTS[2:]
    assert any("2 highlighted" in m for m in findings(tmp_path, {**base, "quadrants": two}, level="ERROR"))
    assert findings(tmp_path, {**base, "quadrants": two, "highlight_set": True}, level="ERROR") == []
    crowded = [{"label": f"i{n}", "cell": 0} for n in range(7)] + [{"label": f"j{n}", "cell": 3} for n in range(14)]
    warns = findings(tmp_path, {**base, "items": crowded}, level="WARN")
    assert any("7 items in matrix cell 0" in m for m in warns) and any("21 matrix items" in m for m in warns)
    errs = findings(tmp_path, {**base, "quadrants": QUADRANTS[:3], "items": [{"label": "x", "cell": 9}]},
                    level="ERROR")
    assert any("needs 4" in m for m in errs) and any("cell from 0 to 3" in m for m in errs)
    flat = {**base, "basis": "measured", "source": "s", "x": {**AXES["x"], "min": 5, "max": 5},
            "thresholds": {"y": [0.3, 0.6]}, "items": [{"label": "a", "x": 5, "y": 0.5}]}
    errs = findings(tmp_path, flat, level="ERROR")
    assert any("x axis max must be greater" in m for m in errs)
    assert any("y threshold needs 1 value" in m for m in errs)
    report = build(deck(tmp_path, flat), tmp_path / "flat.pptx")  # a bad scale must not crash the build
    assert report["slides"] == 1


def test_matrix_three_by_three_and_korean_tag(tmp_path):
    cells = [{"name": f"C{n}"} for n in range(9)]
    report, prs = built(tmp_path, {"layout": "matrix", "size": 3, **AXES, "quadrants": cells,
                                   "items": [{"label": "x", "cell": 8}]}, language="ko")
    assert errors(report) == []
    sh = shapes(prs)
    assert "Cell 9" in sh and "Cell 10" not in sh
    assert sh["Tag"].text_frame.text == "정성적 배치"


# -- symbol tables -----------------------------------------------------------

def angles(shape):
    gds = shape._element.spPr.find(qn("a:prstGeom")).find(qn("a:avLst"))
    return [int(gd.get("fmla").split()[1]) / 60000 for gd in gds]


def test_harvey_cells_are_native_quarters_with_a_legend(tmp_path):
    report, prs = built(tmp_path, {"layout": "table", "symbols": "harvey",
                                   "columns": ["Option", "A", "B", "C", "D", "E", "F"],
                                   "rows": [["x", {"s": 0}, {"s": 1}, {"s": 2}, {"s": 3}, {"s": 4}, {"s": "na"}]]})
    assert errors(report) == []
    sh = list(prs.slides[0].shapes)
    pies = [s for s in sh if s.name.startswith("Symbol harvey") and s.name.endswith("fill")]
    assert [s.auto_shape_type for s in pies] == [MSO_SHAPE.PIE] * 3
    assert [angles(s) for s in pies] == [[270, 0], [270, 90], [270, 180]]  # from 12 o'clock
    full = next(s for s in sh if s.name == "Symbol harvey 4")
    assert full.auto_shape_type == MSO_SHAPE.OVAL and full.fill.type is not None
    assert "Symbol harvey na" in [s.name for s in sh]  # a dash, distinct from an empty circle
    legend = [s.text_frame.text for s in sh if s.name.startswith("Legend text")]
    assert legend == ["None", "Low", "Partial", "High", "Full", "Not assessed"]
    table = next(s for s in sh if s.has_table)
    assert table.table.cell(1, 1).text_frame.text == ""  # the symbol carries the cell
    legend_top = min(s.top for s in sh if s.name.startswith("Legend "))
    assert legend_top >= table.top + table.height


def test_symbol_values_are_normalised_including_yaml_booleans():
    assert normalize_symbol("check", True) == "yes" and normalize_symbol("check", False) == "no"
    assert normalize_symbol("harvey", True) is None  # True is not a quarter count
    assert normalize_symbol("harvey", "3") == 3 and normalize_symbol("tier", "N/A") == "na"
    assert normalize_symbol("rag", "amber") == "a" and normalize_symbol("rag", "blue") is None


def test_rag_cells_always_carry_words_and_check_reads_yaml_yes(tmp_path):
    path = tmp_path / "d.deck.yaml"
    path.write_text("""meta: {audience: a, objective: o, desired_action: d}
slides:
- layout: table
  title: Only one vendor meets every requirement
  notes: n
  symbols: check
  columns: [Requirement, A, B]
  rows:
  - [Sign-on, {s: yes}, {s: no}]
  - [Region, {s: partial}, {s: na}]
- layout: table
  title: Two of four streams are on track today
  notes: n
  symbols: rag
  columns: [Stream, Status]
  rows:
  - [Data, {s: g, label: Approved}]
  - [Change, {s: r}]
""", encoding="utf-8")
    report = build(path, tmp_path / "o.pptx")
    assert errors(report) == []
    prs = Presentation(str(tmp_path / "o.pptx"))
    names = [s.name for s in prs.slides[0].shapes]
    assert "Symbol check yes" in names and "Symbol check no" in names and "Symbol check partial" in names
    cells = texts(prs, 1)
    assert "Approved" in cells and "Off track" in cells  # missing label falls back to the legend word
    warns = [f["message"] for f in lint(path) if f["level"] == "WARN"]
    assert any("status without a label" in m for m in warns)


def test_symbol_lint_blocks_blank_and_invalid_cells(tmp_path):
    t = {"layout": "table", "symbols": "harvey", "columns": ["Option", "Cost", "Speed"]}
    errs = findings(tmp_path, {**t, "rows": [["a", {"s": 2}, {"s": 7}], ["b", "", {"s": 1}]]}, level="ERROR")
    assert any("is empty" in m for m in errs) and any("'7' is not a harvey value" in m for m in errs)
    errs = findings(tmp_path, {**t, "symbols": None, "rows": [["a", {"s": 2}, {"s": 1}]]}, level="ERROR")
    assert any("no 'symbols' family" in m for m in errs)
    warns = findings(tmp_path, {**t, "legend": False, "rows": [["a", {"s": 2}, {"s": 1}]]}, level="WARN")
    assert any("without a legend" in m for m in warns)


def test_symbol_labels_follow_language_brand_and_slide(tmp_path):
    rows = [["a", {"s": 1}, {"s": "na"}]]
    t = {"layout": "table", "symbols": "tier", "columns": ["x", "y", "z"], "rows": rows}
    _, prs = built(tmp_path, t, language="ko")
    assert "미평가" in texts(prs)
    brand = tmp_path / "brand.yaml"
    brand.write_text("copy:\n  symbol_labels:\n    tier: {1: Leader}\n", encoding="utf-8")
    _, prs = built(tmp_path, {**t, "legend": {3: "Laggard"}}, brand=brand)
    legend = texts(prs)
    assert "Leader" in legend and "Laggard" in legend and "Tier 2" in legend


@pytest.mark.skipif(not (shutil.which("soffice") and shutil.which("pdftoppm")),
                    reason="needs LibreOffice and Poppler")
def test_harvey_quarters_render_in_libreoffice(tmp_path):
    from PIL import Image
    _, prs = built(tmp_path, {"layout": "table", "symbols": "harvey", "columns": ["x", "q1", "q2", "q3"],
                              "rows": [["a", {"s": 1}, {"s": 2}, {"s": 3}]]})
    marks = {s.name: s for s in prs.slides[0].shapes if s.name in
             ("Symbol harvey 1", "Symbol harvey 2", "Symbol harvey 3")}
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp_path}/profile", "--headless",
                    "--convert-to", "pdf", "--outdir", str(tmp_path), str(tmp_path / "o.pptx")],
                   check=True, capture_output=True, timeout=180)
    dpi = 288
    subprocess.run(["pdftoppm", "-r", str(dpi), "-png", "-singlefile", str(tmp_path / "o.pdf"),
                    str(tmp_path / "page")], check=True)
    img = Image.open(tmp_path / "page.png").convert("L")

    def dark(shape, dx, dy):  # sample a quarter of the circle, offsets in units of the radius
        cx, cy = centre(shape)
        r = shape.width / PT / 2
        return img.getpixel((round((cx + dx * r * 0.5) * dpi / 72), round((cy + dy * r * 0.5) * dpi / 72))) < 128

    quarters = [(1, -1), (1, 1), (-1, 1), (-1, -1)]  # top-right, bottom-right, bottom-left, top-left
    for n in (1, 2, 3):
        assert [dark(marks[f"Symbol harvey {n}"], *q) for q in quarters] == [i < n for i in range(4)], n


# -- shift -------------------------------------------------------------------

def test_shift_rows_pair_from_and_to(tmp_path):
    report, prs = built(tmp_path, {"layout": "shift", "headers": ["Today", "Target"], "rows": [
        {"dimension": "Intake", "from": "By **email**", "to": "Through **one form**"},
        {"dimension": "Triage", "from": "A person sorts", "to": "Software sorts", "highlight": True}]})
    assert errors(report) == []
    sh = shapes(prs)
    assert sh["From 1"].top == sh["To 1"].top and sh["To 1"].left > sh["Shift arrow 1"].left > sh["From 1"].left
    assert sh["Shift arrow 1"].auto_shape_type == MSO_SHAPE.CHEVRON
    assert str(sh["Shift arrow 2"].fill.fore_color.rgb) == load_brand().hex("accent")
    assert sh["Shift header 1"].text_frame.text == "TODAY"


def test_shift_lint(tmp_path):
    rows = [{"dimension": f"d{n}", "from": "Annual plan", "to": "Quarterly plan"} for n in range(7)]
    rows[0] = {"dimension": "x", "from": "Short", "to": ""}
    rows[1] = {"dimension": "y", "from": "Plan", "to": "A much longer target state description"}
    out = findings(tmp_path, {"layout": "shift", "rows": rows})
    assert any("7 shift rows" in m for m in out)
    assert any("row 1 has no 'to'" in m for m in out)
    assert any("row 2: from and to differ" in m for m in out)


# -- process -----------------------------------------------------------------

STEPS = [{"title": "Prepare", "period": "Q1"}, {"title": "Pilot", "period": "Q2"},
         {"title": "Decide", "period": "Q3"}]
ROWS = [{"label": "Objective", "cells": ["Data", "Live", "Scale"]}]


def test_process_chevron_with_periods_current_and_rows(tmp_path):
    report, prs = built(tmp_path, {"layout": "process", "style": "chevron", "current": 1,
                                   "steps": STEPS, "rows": ROWS})
    assert errors(report) == []
    sh = shapes(prs)
    assert sh["Chevron 1"].auto_shape_type == MSO_SHAPE.PENTAGON
    assert sh["Chevron 2"].auto_shape_type == MSO_SHAPE.CHEVRON
    assert sh["Period 2"].text_frame.text == "Q2 · Now" and sh["Period 1"].text_frame.text == "Q1"
    assert str(sh["Chevron 2"].fill.fore_color.rgb) == load_brand().hex("accent")
    cells = [s for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text in ("Data", "Live", "Scale")]
    for cell, n in zip(cells, (1, 2, 3)):  # rows align under the chevron titles
        assert cell.left == sh[f"Chevron {n} title"].left
    assert sh["Current column"].left >= sh["Chevron 2"].left
    assert sh["Row label"].text_frame.text == "Objective"


def test_process_rail_stays_the_default(tmp_path):
    _, prs = built(tmp_path, {"layout": "process", "steps": ["A", "B", "C"], "current": 0})
    names = [s.name for s in prs.slides[0].shapes]
    assert not any(n.startswith("Chevron") for n in names)
    assert "01" in texts(prs) and "Now" not in " ".join(texts(prs))
    _, prs = built(tmp_path, {"layout": "process", "steps": STEPS, "rows": ROWS})
    assert "01  ·  Q1" in texts(prs) and "Objective" in texts(prs)


def test_process_lint(tmp_path):
    six = [{"title": f"S{n}"} for n in range(7)]
    out = findings(tmp_path, {"layout": "process", "style": "chevron", "steps": six})
    assert any("7 process steps; keep at most 6" in m for m in out)
    out = findings(tmp_path, {"layout": "process", "steps": six[:6]})
    assert any("6 process steps; keep at most 5" in m for m in out)
    out = findings(tmp_path, {"layout": "process", "steps": STEPS, "current": 3,
                              "rows": [{"label": "x", "cells": ["a"]}]})
    assert any("1 cells for 3 steps" in m for m in out) and any("current 3" in m for m in out)


def test_chevron_title_wrap_is_reported(tmp_path):
    report, _ = built(tmp_path, {"layout": "process", "style": "chevron", "steps": [
        {"title": "A very long chevron title that cannot fit"}] + STEPS[1:] * 2})
    assert any("chevron title" in w["message"] for w in report["warnings"])


# -- cards -------------------------------------------------------------------

CARDS = [{"heading": f"Do thing {n}", "body": "Short body"} for n in range(5)]


def test_open_numbered_cards_have_no_fills(tmp_path):
    report, prs = built(tmp_path, {"layout": "cards", "style": "open", "numbered": True,
                                   "cards": CARDS, "highlight": 0})
    assert report["warnings"] == []
    sh = shapes(prs)
    assert [sh[f"Card {n} number"].text_frame.text for n in range(1, 6)] == ["01", "02", "03", "04", "05"]
    assert sh["Card 2"].shape_type == 9  # a hairline, not a panel
    assert sh["Card 1"].height / PT == pytest.approx(2)  # accent rule marks the highlight
    assert sh["Card 1 number"].text_frame.paragraphs[0].runs[0].font.size.pt == load_brand()["type.statement"]
    assert sh["Card 1"].left == sh["Card 1 number"].left  # open text aligns with the column edge


def test_cards_limits(tmp_path):
    report, _ = built(tmp_path, {"layout": "cards", "cards": CARDS})
    assert any("5 cards" in w["message"] for w in report["warnings"])
    assert any("more than 4 cards" in m for m in findings(tmp_path, {"layout": "cards", "cards": CARDS}))
    assert not any("cards" in m for m in findings(tmp_path, {"layout": "cards", "style": "open", "cards": CARDS}))
    uneven = [{"heading": "Fix it"}, {"heading": "Agree the measures with every team before the pilot"}]
    assert any("same grammar" in m for m in findings(tmp_path, {"layout": "cards", "numbered": True,
                                                                "cards": uneven}))


# -- stats -------------------------------------------------------------------

def stat_items(n):
    return [{"value": f"{k}%", "label": f"label {k}"} for k in range(n)]


@pytest.mark.parametrize("n,rows,size", [(5, 2, "statement"), (6, 2, "statement"), (9, 3, "section")])
def test_stats_grid(tmp_path, n, rows, size):
    report, prs = built(tmp_path, {"layout": "stats", "stats": stat_items(n), "source": "s"})
    assert errors(report) == []
    sh = shapes(prs)
    tops = sorted({sh[f"Stat {k}"].top for k in range(1, n + 1)})
    lefts = sorted({sh[f"Stat {k}"].left for k in range(1, n + 1)})
    assert len(tops) == rows and len(lefts) == 3
    assert sh["Stat 4"].left == sh["Stat 1"].left and sh["Stat 4"].top > sh["Stat 1"].top
    pt = sh["Stat 1"].text_frame.paragraphs[0].runs[0].font.size.pt
    assert pt == load_brand()[f"type.{size}"] and pt >= 2.4 * load_brand()["type.body"]
    assert not [s for s in prs.slides[0].shapes if s.shape_type == 1 and s.name.startswith("Stat")]  # no boxes


def test_stats_up_to_four_keep_one_row_and_ten_warn(tmp_path):
    _, prs = built(tmp_path, {"layout": "stats", "stats": stat_items(4), "source": "s"})
    sh = shapes(prs)
    assert len({sh[f"Stat {k}"].top for k in range(1, 5)}) == 1
    report, _ = built(tmp_path, {"layout": "stats", "stats": stat_items(10), "source": "s"})
    assert any("10 stats" in w["message"] for w in report["warnings"])
    assert any("10 stats" in m for m in findings(tmp_path, {"layout": "stats", "stats": stat_items(10),
                                                            "source": "s"}))


# -- example deck --------------------------------------------------------------

def test_example_deck_exercises_every_framework_feature():
    spec = yaml.safe_load((SKILL / "examples/decks/frameworks.deck.yaml").read_text(encoding="utf-8"))
    slides = spec["slides"]
    assert {s.get("symbols") for s in slides if s["layout"] == "table"} >= {"harvey", "check", "rag", "tier"}
    assert {s.get("basis", "qualitative") for s in slides if s["layout"] == "matrix"} == {"qualitative", "measured"}
    assert any(s["layout"] == "matrix" and s.get("size") == 3 for s in slides)
    assert any(s["layout"] == "shift" for s in slides)
    assert any(s["layout"] == "process" and s.get("style") == "chevron" and s.get("rows") for s in slides)
    assert any(s["layout"] == "cards" and s.get("style") == "open" and s.get("numbered") for s in slides)
    assert {len(s["stats"]) for s in slides if s["layout"] == "stats"} >= {6, 9}
