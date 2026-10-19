"""Drawing primitives bound to a brand.

All coordinates are points on the brand canvas. Colors are brand tokens
(``accent``, ``muted`` ...) or literal hex. Fonts are roles (``heading``,
``body``, ``mono``). Use these instead of raw python-pptx calls in custom
slides so every value stays inside the brand profile.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from .brand import Brand, contrast
from . import exhibits, measure

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT,
         "justify": PP_ALIGN.JUSTIFY}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
BOLD = re.compile(r"(\*\*.+?\*\*)")
SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE,
          "oval": MSO_SHAPE.OVAL, "can": MSO_SHAPE.CAN, "chevron": MSO_SHAPE.CHEVRON,
          "pentagon": MSO_SHAPE.PENTAGON, "diamond": MSO_SHAPE.DIAMOND}
CHARTS = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED, "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS, "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "pie": XL_CHART_TYPE.PIE,
    "waterfall": XL_CHART_TYPE.COLUMN_STACKED,  # invisible base series under the steps
}
# Removes PowerPoint's default table style so rules come only from the brand.
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


@dataclass
class Warnings:
    items: list = field(default_factory=list)

    def add(self, slide: int, message: str, level: str = "WARN"):
        self.items.append({"slide": slide, "level": level, "message": message})


class Canvas:
    def __init__(self, slide, brand: Brand, number: int = 0, warnings: Warnings | None = None,
                 dark: bool = False):
        self.slide = slide
        self.brand = brand
        self.number = number
        self.warnings = warnings or Warnings()
        self.dark = dark

    # -- helpers ----------------------------------------------------------
    def warn(self, message: str, level: str = "WARN"):
        self.warnings.add(self.number, message, level)

    @property
    def ink(self) -> str:
        return "on_dark" if self.dark else "ink"

    @property
    def body_color(self) -> str:
        return "on_dark_body" if self.dark else "body"

    @property
    def muted(self) -> str:
        return "on_dark_body" if self.dark else "muted"

    def size(self, role_or_pt) -> float:
        if isinstance(role_or_pt, (int, float)):
            pt = float(role_or_pt)
        else:
            pt = float(self.brand[f"type.{role_or_pt}"])
        minimum = float(self.brand.get("type.min", 0))
        if pt < minimum:
            self.warn(f"text {pt}pt is below type.min {minimum}pt; raised to the minimum")
            pt = minimum
        return pt

    def _font(self, run, role: str, size: float, color: str, bold: bool, italic: bool = False):
        fonts = self.brand.get(f"fonts.{role}") or self.brand["fonts.body"]
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = self.brand.rgb(color)
        rpr = run._r.get_or_add_rPr()
        for tag, face in (("a:latin", fonts.get("latin")), ("a:ea", fonts.get("east_asian")),
                          ("a:cs", fonts.get("latin"))):
            if not face:
                continue
            el = rpr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rpr, qn(tag))
            el.set("typeface", face)
        # schema order inside rPr: ln, fill, effect, ..., latin, ea, cs
        for tag in ("a:latin", "a:ea", "a:cs"):
            el = rpr.find(qn(tag))
            if el is not None:
                rpr.remove(el)
                rpr.append(el)

    @staticmethod
    def _plain(shape):
        """Drop the theme style reference so no inherited shadow or effect applies."""
        style = shape._element.find(qn("p:style"))
        if style is not None:
            shape._element.remove(style)
        return shape

    @staticmethod
    def _no_inset(tf, inset: float | tuple = 0):
        l, t, r, b = inset if isinstance(inset, tuple) else (inset,) * 4
        tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (Pt(l), Pt(t), Pt(r), Pt(b))

    # -- text ---------------------------------------------------------------
    def text(self, x, y, w, h, content, *, size="body", color=None, font="body", bold=False,
             italic=False, align="left", anchor="top", spacing=None, inset=0, wrap=True,
             bold_color=None, bullet=None, paragraph_gap=0.0, name=None, check_fit=True,
             fill=None):
        """Add a text box.

        ``content`` is a string (``\\n`` separates paragraphs, ``**x**`` is bold)
        or a list of strings (one paragraph each).
        ``bullet`` is a character such as "•" or "–" to prefix every paragraph
        with a hanging indent.
        """
        pt = self.size(size)
        color = color or self.body_color
        spacing = spacing or float(self.brand.get("type.line_spacing", 1.2))
        paras = content if isinstance(content, list) else str(content).split("\n")
        box = self.slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        if name:
            box.name = name
        if fill:  # e.g. masks the line behind an edge label
            box.fill.solid()
            box.fill.fore_color.rgb = self.brand.rgb(fill)
        tf = box.text_frame
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE  # boxes keep their planned size; QA catches overflow
        self._no_inset(tf, inset)
        tf.vertical_anchor = ANCHOR[anchor]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN[align]
            p.line_spacing = spacing
            if paragraph_gap and i > 0:
                p.space_before = Pt(paragraph_gap)
            if bullet:
                self._bullet(p, bullet, pt)
            for piece in BOLD.split(str(para)):
                if not piece:
                    continue
                strong = piece.startswith("**") and piece.endswith("**")
                run = p.add_run()
                run.text = piece[2:-2] if strong else piece
                self._font(run, font, pt, (bold_color or color) if strong else color,
                           bold or strong, italic)
        if check_fit:
            inner_w = w - (inset[0] + inset[2] if isinstance(inset, tuple) else 2 * inset)
            need = sum(measure.text_height(p, pt, inner_w - (pt if bullet else 0), spacing, font)
                       for p in paras) + paragraph_gap * (len(paras) - 1)
            if need > h + 0.5:
                label = measure.strip_markup(paras[0])[:40]
                self.warn(f"text may overflow its box ({need:.0f}pt needed, {h:.0f}pt available): '{label}'")
        return box

    def _bullet(self, p, char: str, pt: float, left: float = 0.0):
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", str(int(Pt(left + pt * 1.1))))
        ppr.set("indent", str(-int(Pt(pt * 1.1))))
        for tag in ("a:buNone", "a:buChar", "a:buClr"):
            for el in ppr.findall(qn(tag)):
                ppr.remove(el)
        clr = etree.SubElement(ppr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr")).set("val", self.brand.hex("accent"))
        etree.SubElement(ppr, qn("a:buChar")).set("char", char)

    def paragraphs(self, x, y, w, paras, *, name=None):
        """One text box holding several differently styled paragraphs.

        ``paras`` is a list of dicts: ``text``, ``size``, ``color``, ``bold_color``,
        ``bullet`` (character or None), ``left`` (extra indent in pt) and
        ``space_before`` (pt). The application flows the lines, so the gaps between
        paragraphs stay even whatever the line breaks turn out to be. Returns the
        box and its estimated height.
        """
        spacing = float(self.brand.get("type.line_spacing", 1.2))
        height = 0.0
        for k, para in enumerate(paras):
            pt = self.size(para.get("size", "body"))
            inner = w - para.get("left", 0.0) - (pt * 1.1 if para.get("bullet") else 0)
            height += measure.text_height(para["text"], pt, inner, spacing)
            height += para.get("space_before", 0.0) if k else 0.0
        box = self.slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(height))
        if name:
            box.name = name
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        self._no_inset(tf, 0)
        for k, para in enumerate(paras):
            pt = self.size(para.get("size", "body"))
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = ALIGN["left"]
            p.line_spacing = spacing
            if k and para.get("space_before"):
                p.space_before = Pt(para["space_before"])
            if para.get("bullet"):
                self._bullet(p, para["bullet"], pt, para.get("left", 0.0))
            color = para.get("color") or self.body_color
            for piece in BOLD.split(str(para["text"])):
                if not piece:
                    continue
                strong = piece.startswith("**") and piece.endswith("**")
                run = p.add_run()
                run.text = piece[2:-2] if strong else piece
                self._font(run, "body", pt, (para.get("bold_color") or color) if strong else color, strong)
        return box, height

    def text_height(self, content, size, width, spacing=None, font="body") -> float:
        spacing = spacing or float(self.brand.get("type.line_spacing", 1.2))
        paras = content if isinstance(content, list) else str(content).split("\n")
        return sum(measure.text_height(p, self.size(size), width, spacing, font) for p in paras)

    # -- shapes -------------------------------------------------------------
    def box(self, x, y, w, h, *, fill=None, line=None, line_width=0.75, dash=False,
            shape="rect", radius=None, name=None, text=None, text_size="body",
            text_color=None, align="center", anchor="middle", inset=6, bold=False):
        radius = self.brand.get("layout.radius", 0) if radius is None else radius
        kind = SHAPES["round"] if (shape == "rect" and radius) else SHAPES[shape]
        shp = self.slide.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
        if name:
            shp.name = name
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = min(0.5, radius / max(1.0, min(w, h)))
        self._plain(shp)
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = self.brand.rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = self.brand.rgb(line)
            shp.line.width = Pt(line_width)
            if dash:
                self._dash(shp.line._get_or_add_ln())
        else:
            shp.line.fill.background()
        tf = shp.text_frame
        tf.word_wrap = True
        self._no_inset(tf, inset)
        tf.vertical_anchor = ANCHOR[anchor]
        if text is not None:
            pt = self.size(text_size)
            for i, para in enumerate(str(text).split("\n")):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = ALIGN[align]
                p.line_spacing = float(self.brand.get("type.heading_line_spacing", 1.05))
                run = p.add_run()
                run.text = para
                self._font(run, "body", pt, text_color or self.ink, bold)
        return shp

    @staticmethod
    def _dash(ln):
        for el in ln.findall(qn("a:prstDash")):
            ln.remove(el)
        dash = etree.SubElement(ln, qn("a:prstDash"))
        dash.set("val", "dash")
        # prstDash must precede round/bevel/miter, headEnd, tailEnd
        for tag in ("a:round", "a:bevel", "a:miter", "a:headEnd", "a:tailEnd"):
            for el in ln.findall(qn(tag)):
                ln.remove(el)
                ln.append(el)

    @staticmethod
    def _arrow(ln, end: str = "tail", kind: str = "triangle"):
        tag = qn(f"a:{end}End")
        for el in ln.findall(tag):
            ln.remove(el)
        el = etree.SubElement(ln, tag)
        el.set("type", kind)
        el.set("w", "med")
        el.set("len", "med")
        if end == "head":  # headEnd must precede tailEnd
            tail = ln.find(qn("a:tailEnd"))
            if tail is not None:
                ln.remove(tail)
                ln.append(tail)

    def line(self, x1, y1, x2, y2, *, color="line", width=0.75, dash=False, arrow=None,
             name=None):
        """Straight line. ``arrow``: None | 'end' | 'both'."""
        conn = self._plain(self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2)))
        if name:
            conn.name = name
        conn.line.color.rgb = self.brand.rgb(color)
        conn.line.width = Pt(width)
        ln = conn.line._get_or_add_ln()
        if dash:
            self._dash(ln)
        if arrow in ("end", "both"):
            self._arrow(ln, "tail")
        if arrow == "both":
            self._arrow(ln, "head")
        return conn

    def connect(self, a, b, *, a_side="right", b_side="left", elbow=True, color="muted",
                width=1.0, dash=False, arrow="end", name=None):
        """Glued connector between two shapes; it follows them when edited."""
        sites = {"top": 0, "left": 1, "bottom": 2, "right": 3}
        kind = MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT
        conn = self._plain(self.slide.shapes.add_connector(kind, 0, 0, 0, 0))
        conn.begin_connect(a, sites[a_side])
        conn.end_connect(b, sites[b_side])
        if name:
            conn.name = name
        conn.line.color.rgb = self.brand.rgb(color)
        conn.line.width = Pt(width)
        ln = conn.line._get_or_add_ln()
        if dash:
            self._dash(ln)
        if arrow in ("end", "both"):
            self._arrow(ln, "tail")
        if arrow == "both":
            self._arrow(ln, "head")
        return conn

    def path(self, points, *, color="muted", width=1.0, dash=False, arrow="end", name=None):
        """Open polyline through ``points`` [(x, y), ...] — for routes a straight or
        elbow connector cannot draw (e.g. a return path that loops under its nodes).
        Not glued: re-check it after moving the nodes it connects."""
        (x0, y0), rest = points[0], points[1:]
        fb = self.slide.shapes.build_freeform(Pt(x0), Pt(y0), scale=1.0)
        fb.add_line_segments([(Pt(x), Pt(y)) for x, y in rest], close=False)
        shp = self._plain(fb.convert_to_shape())
        if name:
            shp.name = name
        shp.fill.background()
        shp.line.color.rgb = self.brand.rgb(color)
        shp.line.width = Pt(width)
        ln = shp.line._get_or_add_ln()
        if dash:
            self._dash(ln)
        if arrow in ("end", "both"):
            self._arrow(ln, "tail")
        if arrow == "both":
            self._arrow(ln, "head")
        return shp

    # -- images -------------------------------------------------------------
    def image(self, path, x, y, w, h, *, fit="contain", name=None, alt=""):
        path = Path(path)
        if not path.exists():
            self.warn(f"image not found: {path}", "ERROR")
            return self.box(x, y, w, h, line="chrome", dash=True, text=f"missing: {path.name}",
                            text_size="caption", text_color="muted")
        if path.suffix.lower() == ".svg":
            png = svg_to_png(path)
            if png is None:
                self.warn(f"cannot place SVG without a converter (install cairosvg or rsvg-convert): {path.name}", "ERROR")
                return None
            path = png
        from PIL import Image
        with Image.open(path) as im:
            iw, ih = im.size
        scale = (min if fit == "contain" else max)(w / iw, h / ih)
        pw, ph = iw * scale, ih * scale
        px, py = x + (w - pw) / 2, y + (h - ph) / 2
        pic = self.slide.shapes.add_picture(str(path), Pt(px), Pt(py), Pt(pw), Pt(ph))
        if fit == "cover":  # crop to the frame instead of overflowing it
            cx = max(0.0, (pw - w) / pw / 2)
            cy = max(0.0, (ph - h) / ph / 2)
            pic.crop_left = pic.crop_right = cx
            pic.crop_top = pic.crop_bottom = cy
            pic.left, pic.top, pic.width, pic.height = Pt(x), Pt(y), Pt(w), Pt(h)
        if name:
            pic.name = name
        if alt:
            pic._element.nvPicPr.cNvPr.set("descr", alt)
        return pic

    # -- table --------------------------------------------------------------
    def table(self, x, y, w, columns, rows, *, widths=None, numeric_cols=(), highlight_col=None,
              total_row=False, size="body", row_h=None, name=None, symbols=None, symbol_labels=None):
        """Quiet native table: thin horizontal rules, no vertical lines.

        A cell may be a string or ``[keyword, detail]`` (bold keyword over a
        muted detail line). With ``symbols`` (``harvey``, ``check``, ``rag``,
        ``tier``) a cell may also be ``{s: value, label: text}``: a native symbol
        centred in the cell, or left of its label. ``symbol_labels`` maps values
        to the words used when a status cell has no label of its own.
        """
        from .frameworks import draw_symbol, normalize_symbol, symbol_size
        n_cols = len(columns)
        widths = widths or [1] * n_cols
        total = float(sum(widths))
        col_w = [w * v / total for v in widths]
        pt = self.size(size)
        detail_pt = self.size("caption")
        pad = 5.0
        sym_d = symbol_size(self, size) if symbols else 0.0
        sym_cols = {c for row in rows for c, cell in enumerate(row) if isinstance(cell, dict) and "s" in cell}

        def symbol_cell(cell):
            """(value, label) of a symbol cell; status cells always carry a word."""
            value = normalize_symbol(symbols, cell.get("s"))
            if value is None:
                self.warn(f"'{cell.get('s')}' is not a {symbols} symbol value", "ERROR")
            label = cell.get("label")
            if label is None and symbols == "rag":
                label = (symbol_labels or {}).get(value, "")
            return value, str(label or "")

        parsed = {}
        heights = [pt * 1.8 + 2 * pad]
        for r, row in enumerate(rows, start=1):
            hmax = 0.0
            for c, (cell, cw) in enumerate(zip(row, col_w)):
                if symbols and isinstance(cell, dict) and "s" in cell:
                    parsed[r, c] = symbol_cell(cell)
                    label = parsed[r, c][1]
                    hh = max(sym_d, self.text_height(label, pt, cw - 2 * pad - sym_d - 6) if label else 0)
                    hmax = max(hmax, hh)
                    continue
                head, detail = (cell + [""])[:2] if isinstance(cell, list) else (cell, "")
                hh = self.text_height(str(head), pt, cw - 2 * pad)
                if detail:
                    hh += self.text_height(str(detail), detail_pt, cw - 2 * pad)
                hmax = max(hmax, hh)
            heights.append(max(row_h or 0, hmax + 2 * pad))
        gf = self.slide.shapes.add_table(len(rows) + 1, n_cols, Pt(x), Pt(y), Pt(w), Pt(sum(heights)))
        if name:
            gf.name = name
        tbl = gf.table
        tbl_pr = tbl._tbl.tblPr
        tbl_pr.set("firstRow", "1")
        tbl_pr.set("bandRow", "0")
        style = tbl_pr.find(qn("a:tableStyleId"))
        if style is None:
            style = etree.SubElement(tbl_pr, qn("a:tableStyleId"))
        style.text = NO_STYLE_TABLE
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Pt(cw)
        for r, hh in enumerate(heights):
            tbl.rows[r].height = Pt(hh)
        tint = self.brand.tint("accent", 0.93)
        last = len(rows)
        glyphs = []
        for r in range(last + 1):
            for c in range(n_cols):
                cell = tbl.cell(r, c)
                cell.margin_left = cell.margin_right = Pt(pad)
                cell.margin_top = cell.margin_bottom = Pt(pad)
                cell.vertical_anchor = MSO_ANCHOR.BOTTOM if r == 0 else MSO_ANCHOR.TOP
                highlighted = highlight_col is not None and c == highlight_col
                if highlighted and r > 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = self.brand.rgb(tint)
                else:
                    cell.fill.background()
                value = columns[c] if r == 0 else rows[r - 1][c] if c < len(rows[r - 1]) else ""
                if symbols and r > 0 and isinstance(value, dict) and "s" in value:
                    sym, label = parsed[r, c]
                    cx0 = x + sum(col_w[:c])
                    first_line = pt * 1.2 * float(self.brand.get("type.line_spacing", 1.2))
                    line_mid = y + sum(heights[:r]) + pad + first_line / 2
                    if label:  # the symbol leads, the word follows it
                        cell.margin_left = Pt(pad + sym_d + 6)
                        glyphs.append((sym, cx0 + pad + sym_d / 2, line_mid))
                    else:
                        glyphs.append((sym, cx0 + col_w[c] / 2, line_mid))
                    value = label
                head, detail = (value + [""])[:2] if isinstance(value, list) else (value, "")
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT if c in numeric_cols else PP_ALIGN.LEFT
                if symbols and c in sym_cols and r == 0 and not any(
                        isinstance(rw[c], dict) and rw[c].get("label") for rw in rows if c < len(rw)):
                    p.alignment = PP_ALIGN.CENTER  # header sits over centred symbols
                if str(head):
                    run = p.add_run()
                    run.text = str(head)
                    if r == 0:
                        self._font(run, "body", self.size("caption"),
                                   "accent" if highlighted else "muted", highlighted)
                    else:
                        strong = bool(detail) and c > 0 or (total_row and r == last)
                        self._font(run, "body", pt, "ink", strong)
                else:  # no empty run: LibreOffice would size the line at its 18 pt default
                    p._p.get_or_add_endParaRPr().set("sz", str(int(pt * 100)))
                if detail and r > 0:
                    p2 = tf.add_paragraph()
                    p2.alignment = p.alignment
                    run2 = p2.add_run()
                    run2.text = str(detail)
                    self._font(run2, "body", detail_pt, "muted", False)
                # borders: header bottom, row separators, table end
                bottom = ("line", 1.0) if r in (0, last) else ("line", 0.5)
                top = ("muted", 1.0) if total_row and r == last else None
                self._cell_borders(cell, top=top, bottom=bottom)
        if highlight_col is not None:  # accent cap above the recommended column
            cx = x + sum(col_w[:highlight_col])
            self.box(cx, y - 2, col_w[highlight_col], 2, fill="accent")
        for sym, cx, cy in glyphs:  # drawn over the table so they stay centred in their cells
            if sym is not None:
                draw_symbol(self, symbols, sym, cx, cy, sym_d)
        return gf, sum(heights)

    def _cell_borders(self, cell, top=None, bottom=None):
        tc_pr = cell._tc.get_or_add_tcPr()
        for side in ("lnL", "lnR", "lnT", "lnB"):
            for el in tc_pr.findall(qn(f"a:{side}")):
                tc_pr.remove(el)
        spec = {"lnL": None, "lnR": None, "lnT": top, "lnB": bottom}
        # border elements must come first inside tcPr, in L, R, T, B order
        for index, side in enumerate(("lnL", "lnR", "lnT", "lnB")):
            ln = etree.Element(qn(f"a:{side}"))
            value = spec[side]
            if value:
                color, width = value
                ln.set("w", str(int(Pt(width))))
                fill = etree.SubElement(ln, qn("a:solidFill"))
                etree.SubElement(fill, qn("a:srgbClr")).set("val", self.brand.hex(color))
            else:
                ln.set("w", "0")
                etree.SubElement(ln, qn("a:noFill"))
            tc_pr.insert(index, ln)

    # -- chart --------------------------------------------------------------
    def chart(self, x, y, w, h, kind, categories, series, *, highlight=None, number_format=None,
              labels=True, legend=None, gap=60, name=None, highlight_series=None, series_kind=None,
              end_labels=None, estimates=None, annotations=None, vmin=None, vmax=None,
              value_range=None, min_pads=None):
        """Native chart styled from the brand.

        ``series`` is ``[{"name": str, "values": [...]}, ...]``. With one series,
        ``highlight`` is a category index (or list) drawn in the accent while the
        others stay muted. ``highlight_series`` (index or list) does the same for
        whole series. Otherwise ``series_kind`` picks the colors: ``ordered`` is a
        ramp of accent tints from dark to light, ``nominal`` is ``colors.series``;
        the default comes from ``colors.series_mode``.

        ``end_labels`` (line charts; default with up to 6 series) writes the series
        name and last value at each line's right end instead of a legend.
        ``estimates`` (``{from: category}``) dashes lines and lightens bars from that
        category on. ``annotations`` adds reference lines, bands, callouts and
        computed deltas. Those need a known plot rectangle, so the plot area is
        then written as a manual layout and the value axis pinned; ``vmin``/``vmax``
        override the range and ``value_range``/``min_pads`` let panels share one
        scale and baseline. ``labels=None`` means on unless end labels are drawn.
        With fixed geometry the returned frame carries ``plot`` (an ``exhibits.Plot``)
        for custom renderers that draw more on the data.
        """
        b = self.brand
        cats = list(categories)
        n_series = len(series)
        single = n_series == 1
        horizontal = kind in exhibits.HORIZONTAL
        est = exhibits.estimate_start(cats, estimates)
        suffix = (estimates or {}).get("suffix")
        display = [f"{cat}{suffix}" if suffix and est is not None and j >= est else cat
                   for j, cat in enumerate(cats)]
        labels, end_labels = exhibits.label_modes(kind, n_series, labels, end_labels)
        if legend is None:
            legend = not single and not end_labels
        annotations = list(annotations or [])
        if annotations and kind not in exhibits.ANNOTATED:
            self.warn(f"annotations work on column, bar and line charts; ignored on '{kind}'")
            annotations = []
        fixed = kind in exhibits.ANNOTATED and (bool(annotations or end_labels or value_range or min_pads)
                                                or vmin is not None or vmax is not None)
        colors = exhibits.series_colors(b, kind, n_series, series_kind, highlight_series)
        focus = set(exhibits.as_list(highlight_series))
        marked = {exhibits.category_index(cats, ref) for ref in exhibits.as_list(highlight)} - {None}
        context = b.tint("muted", float((b.get("chart.context_tints") or [0.45])[0]))
        ends = exhibits.end_texts(series, labels, number_format) if end_labels else []
        plan = geom = None
        if fixed:
            plan = exhibits.plan_chart(self, kind, w, h, display, series, labels=labels,
                                       number_format=number_format, legend=bool(legend),
                                       end_texts=[t for _, _, t in ends], annotations=annotations,
                                       vmin=vmin, vmax=vmax, value_range=value_range,
                                       min_pads=min_pads, gap=gap)
            geom = plan.plot(x, y, w, h, len(cats), horizontal)
            exhibits.draw_bands(self, geom, plan, cats, annotations)  # behind the chart
        # An estimated stretch of a line is its own dashed series from the last actual
        # point on: renderers that ignore per-point line styles still show the dash.
        split = kind == "line" and bool(est)
        data = CategoryChartData()
        data.categories = display
        for s in series:
            data.add_series(s["name"], s["values"][:est] + [None] * (len(cats) - est) if split else s["values"])
        if split:
            label = exhibits.words(b, getattr(self, "language", "en"))["estimate_series"]
            for s in series:
                data.add_series(label.format(s["name"]), [None] * (est - 1) + list(s["values"][est - 1:]))
        frame = self.slide.shapes.add_chart(CHARTS[kind], Pt(x), Pt(y), Pt(w), Pt(h), data)
        if name:
            frame.name = name
        chart = frame.chart
        chart.has_title = False  # the slide headline carries the message
        chart.font.size = Pt(self.size("caption"))
        chart.font.color.rgb = self.brand.rgb(self.muted)
        fonts = self.brand["fonts.body"]
        chart.font.name = fonts.get("latin")
        chart.has_legend = bool(legend)
        if legend:
            chart.legend.position = XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(self.size("caption"))
        plot = chart.plots[0]
        if kind in ("bar", "column", "stacked_bar", "stacked_column"):
            plot.gap_width = gap
            if kind.startswith("stacked"):
                plot.overlap = 100
        if kind not in ("pie", "doughnut"):
            va = chart.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = self.brand.rgb("line")
            va.major_gridlines.format.line.width = Pt(0.5)
            va.format.line.fill.background()
            va.tick_labels.font.size = Pt(self.size("caption"))
            if kind in ("bar", "column", "stacked_bar", "stacked_column"):
                va.minimum_scale = 0  # bars encode length; they always start at zero
            if labels:
                va.visible = False
                va.has_major_gridlines = False
            elif number_format:
                va.tick_labels.number_format = number_format
                va.tick_labels.number_format_is_linked = False
            ca = chart.category_axis
            ca.format.line.color.rgb = self.brand.rgb("line")
            ca.tick_labels.font.size = Pt(self.size("caption"))
            ca.tick_labels.font.color.rgb = self.brand.rgb(self.body_color)
            if kind in ("bar", "stacked_bar"):
                ca.reverse_order = True  # first category on top, as read

        def point_color(i, j):
            if single and marked and kind not in ("line",):
                return "accent" if j in marked else context
            return colors[i]

        for i, s in enumerate(plot.series):
            fmt = s.format
            if kind == "line":
                k = i % n_series  # estimate series follow the actual ones
                lead = k in focus if focus else k == 0
                fmt.line.color.rgb = b.rgb(colors[k])
                fmt.line.width = Pt(2.25 if lead else 1.5)
                if i >= n_series or est == 0:
                    fmt.line.dash_style = MSO_LINE_DASH_STYLE.DASH
                s.smooth = False
                s.marker.style = XL_MARKER_STYLE.CIRCLE
                s.marker.size = 5
                s.marker.format.fill.solid()
                s.marker.format.fill.fore_color.rgb = b.rgb(colors[k])
                s.marker.format.line.color.rgb = b.rgb(colors[k])
                continue
            fmt.fill.solid()
            fmt.fill.fore_color.rgb = b.rgb(colors[i] if not (single and marked) else context)
            if kind in ("pie", "doughnut"):
                for j in range(len(cats)):
                    pt_ = s.points[j]
                    pt_.format.fill.solid()
                    pt_.format.fill.fore_color.rgb = b.rgb(
                        b.series(j) if not marked else ("accent" if j in marked else context))
                continue
            for j in range(len(cats)):
                estimate = est is not None and j >= est
                if not (estimate or (single and j in marked)):
                    continue
                base = point_color(i, j)
                pfmt = s.points[j].format
                pfmt.fill.solid()
                pfmt.fill.fore_color.rgb = b.rgb(
                    b.tint(base, float(b.get("chart.estimate_tint", 0.55))) if estimate else base)
                if estimate:  # outlined in the actual color so the series stays recognisable
                    pfmt.line.color.rgb = b.rgb(base)
                    pfmt.line.width = Pt(0.75)
        if labels:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(self.size("caption"))
            dl.font.color.rgb = self.brand.rgb(self.ink)
            if number_format:
                dl.number_format = number_format
                dl.number_format_is_linked = False
            if kind in ("bar", "column"):
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            elif kind == "line":
                dl.position = XL_LABEL_POSITION.ABOVE
            for i, s in enumerate(plot.series):
                if split and i >= n_series:  # the joint point is labeled by the actual series
                    self._series_labels(s, self.ink, number_format, XL_LABEL_POSITION.ABOVE, hide=[est - 1])
                elif kind.startswith("stacked"):  # labels sit on the fill: pick the readable text color
                    on = "on_accent" if contrast(b.hex("on_accent"), b.hex(colors[i])) > \
                        contrast(b.hex(self.ink), b.hex(colors[i])) else self.ink
                    if on != self.ink:
                        self._series_labels(s, on, number_format, None)
        if split and legend:
            self._hide_legend_entries(chart, range(n_series, 2 * n_series))
        if fixed:
            self._fix_plot(chart, plan, w, h, see_through=any(a.get("kind") == "band" for a in annotations))
            side = exhibits.draw_ref_lines(self, geom, series, annotations, number_format)
            background = "dark" if self.dark else "surface"
            side += [{"y": geom.pos(value), "text": text, "name": "End label",
                      "color": exhibits.readable(b, colors[i], background, self.muted)}
                     for i, value, text in ends]
            exhibits.side_labels(self, geom, side, y, y + h)
            exhibits.draw_lanes(self, geom, plan, kind, number_format, labels)
            frame.plot = geom  # geometry for callers that draw more on the chart
        return frame

    def _series_labels(self, ser, color, number_format, position, hide=()):
        """Series-level data labels; they replace the chart-level ones for this series."""
        dl = ser.data_labels
        dl.show_value = True
        dl.font.size = Pt(self.size("caption"))
        dl.font.color.rgb = self.brand.rgb(color)
        if number_format:
            dl.number_format = number_format
            dl.number_format_is_linked = False
        if position is not None:
            dl.position = position
        for idx in hide:
            lbl = etree.Element(qn("c:dLbl"))
            etree.SubElement(lbl, qn("c:idx")).set("val", str(idx))
            etree.SubElement(lbl, qn("c:delete")).set("val", "1")
            dl._element.insert(0, lbl)

    @staticmethod
    def _hide_legend_entries(chart, indexes):
        legend = chart._chartSpace.find(qn("c:chart")).find(qn("c:legend"))
        anchor = legend.find(qn("c:legendPos"))
        for idx in sorted(indexes, reverse=True):
            entry = etree.Element(qn("c:legendEntry"))
            etree.SubElement(entry, qn("c:idx")).set("val", str(idx))
            etree.SubElement(entry, qn("c:delete")).set("val", "1")
            anchor.addnext(entry)

    def waterfall(self, x, y, w, h, spec, *, number_format=None, gap=50, name=None,
                  value_range=None, min_pads=None):
        """Bridge from a start value to an end value through signed steps.

        ``spec`` holds ``steps: [{label, value}]`` and optionally ``start: {label,
        value}`` and ``total`` (label of the end bar). Drawn as a stacked column
        whose first series is an invisible base: rises in ``colors.positive``, falls
        in ``colors.negative``, start and end in ``colors.ink``, each bar labeled.
        """
        b = self.brand
        cats, rows, series = exhibits.waterfall_series(spec)
        if exhibits.waterfall_low(spec) < 0:
            self.warn("waterfall running total goes below zero; a stacked bridge cannot show it — "
                      "rebase the start or split the bridge", "ERROR")
        plan = exhibits.plan_chart(self, "waterfall", w, h, cats, series, labels=True,
                                   number_format=number_format, value_range=value_range,
                                   min_pads=min_pads)
        plot = plan.plot(x, y, w, h, len(cats), False)
        data = CategoryChartData()
        data.categories = cats
        for s in series:
            data.add_series(s["name"], s["values"])
        frame = self.slide.shapes.add_chart(CHARTS["waterfall"], Pt(x), Pt(y), Pt(w), Pt(h), data)
        if name:
            frame.name = name
        chart = frame.chart
        chart.has_title = False
        chart.has_legend = False
        chart.font.size = Pt(self.size("caption"))
        chart.font.color.rgb = b.rgb(self.muted)
        chart.font.name = b["fonts.body"].get("latin")
        plot_ = chart.plots[0]
        plot_.gap_width = gap
        plot_.overlap = 100
        fills = {"Total": self.ink, "Increase": "positive", "Decrease": "negative"}
        for s in plot_.series:
            if s.name == "Base":
                s.format.fill.background()
                s.format.line.fill.background()
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = b.rgb(fills[s.name])
        va = chart.value_axis
        va.visible = False
        va.has_major_gridlines = False
        ca = chart.category_axis
        ca.format.line.color.rgb = b.rgb("line")
        ca.tick_labels.font.size = Pt(self.size("caption"))
        ca.tick_labels.font.color.rgb = b.rgb(self.body_color)
        self._fix_plot(chart, plan, w, h)
        cap = self.size("caption")
        for j, (_, base, length, role, value) in enumerate(rows):
            text = exhibits.format_number(value, number_format) if role == "total" \
                else exhibits.signed(value, number_format)
            a, z = plot.slot(j)
            top = plot.pos(base + length)
            self.text(a, top - cap * 1.4, z - a, cap * 1.4, text, size="caption", color=self.ink,
                      align="center", anchor="bottom", wrap=False, check_fit=False, name="Value label")
        frame.plot = plot
        return frame

    @staticmethod
    def _fix_plot(chart, plan, w, h, see_through=False):
        """Write the plot area as a manual inner layout and pin the value axis, so
        shapes drawn at ``Plot.pos`` land on the data in any renderer."""
        space = chart._chartSpace
        plot_area = space.find(qn("c:chart")).find(qn("c:plotArea"))
        for old in plot_area.findall(qn("c:layout")):
            plot_area.remove(old)
        layout = etree.Element(qn("c:layout"))
        manual = etree.SubElement(layout, qn("c:manualLayout"))
        left, top, right, bottom = plan.pads
        for tag, val in (("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"),
                         ("x", left / w), ("y", top / h), ("w", (w - left - right) / w),
                         ("h", (h - top - bottom) / h)):
            etree.SubElement(manual, qn(f"c:{tag}")).set("val", f"{val:.5f}" if isinstance(val, float) else val)
        plot_area.insert(0, layout)
        va = chart.value_axis
        va.minimum_scale = plan.lo
        va.maximum_scale = plan.hi
        if plan.step:
            va.major_unit = plan.step
        val_ax = va._element
        if val_ax.find(qn("c:crossBetween")) is None:  # bars and points sit mid-slot
            crosses = val_ax.find(qn("c:crosses"))
            crosses = crosses if crosses is not None else val_ax.find(qn("c:crossesAt"))
            between = etree.Element(qn("c:crossBetween"))
            between.set("val", "between")
            crosses.addnext(between)
        if see_through:  # bands are drawn behind the chart; keep both areas unfilled
            for parent, after in ((space, space.find(qn("c:chart"))), (plot_area, None)):
                if parent.find(qn("c:spPr")) is not None:
                    parent.remove(parent.find(qn("c:spPr")))
                sp = etree.Element(qn("c:spPr"))
                etree.SubElement(sp, qn("a:noFill"))
                etree.SubElement(etree.SubElement(sp, qn("a:ln")), qn("a:noFill"))
                if after is not None:
                    after.addnext(sp)
                else:
                    ext = parent.find(qn("c:extLst"))
                    (ext.addprevious(sp) if ext is not None else parent.append(sp))


def svg_to_png(path: Path) -> Path | None:
    out = Path(tempfile.mkdtemp(prefix="deckkit-")) / (path.stem + ".png")
    try:
        import cairosvg  # type: ignore
        cairosvg.svg2png(url=str(path), write_to=str(out), output_width=512)
        return out
    except Exception:
        pass
    if shutil.which("rsvg-convert"):
        subprocess.run(["rsvg-convert", "-w", "512", "-o", str(out), str(path)], check=False)
        if out.exists():
            return out
    return None
