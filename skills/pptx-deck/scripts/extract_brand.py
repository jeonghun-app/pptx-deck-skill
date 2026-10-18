#!/usr/bin/env python3
"""Draft a brand profile from an existing company template or deck.

    python scripts/extract_brand.py company-template.pptx -o brand/brand.yaml [--report brand/extract-report.md]

Reads the measurable part of a design: slide size, theme colors and fonts,
master/layout names and placeholder geometry, and (when the file has slides)
the font sizes and colors actually used. The draft keeps the template path so
decks are built on the template's masters. Review every value — the script
cannot know which color is the accent or which layout your team prefers;
record those decisions in brand/guidelines.md.
"""
import argparse
import collections
import shutil
import sys
from pathlib import Path

import yaml
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit.brand import DEFAULT_BRAND, contrast, luminance  # noqa: E402

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
ROLE_HINTS = {
    "cover": ("title slide", "cover", "제목 슬라이드", "표지"),
    "section": ("section", "divider", "구역", "간지", "섹션"),
    "content": ("title only", "제목만", "title and content", "제목 및 내용"),
    "closing": ("closing", "end", "thank", "마무리", "엔딩"),
    "blank": ("blank", "빈 화면"),
}


def pt(emu) -> float:
    return round(Emu(int(emu)).pt, 1)


def theme(prs):
    part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    colors = {}
    scheme = root.find(".//a:clrScheme", NS)
    for el in (scheme if scheme is not None else []):
        name = etree.QName(el).localname
        child = el[0]
        # srgbClr carries val="RRGGBB"; sysClr carries val="windowText" lastClr="000000"
        value = child.get("lastClr") if etree.QName(child).localname == "sysClr" else child.get("val")
        if value and len(value) == 6:
            colors[name] = value
    fonts = {}
    for kind in ("majorFont", "minorFont"):
        node = root.find(f".//a:{kind}", NS)
        if node is None:
            continue
        latin = node.find("a:latin", NS)
        ea = node.find("a:ea", NS)
        hang = node.find("a:font[@script='Hang']", NS)
        fonts[kind] = {
            "latin": latin.get("typeface") if latin is not None else None,
            "east_asian": (ea.get("typeface") if ea is not None and ea.get("typeface") else None)
            or (hang.get("typeface") if hang is not None else None),
        }
    return colors, fonts


def master_sizes(prs):
    """Level-1 title and body sizes (pt) from the first master's text styles."""
    el = prs.slide_masters[0]._element
    out = {}
    for kind in ("titleStyle", "bodyStyle"):
        node = el.find(f".//{{http://schemas.openxmlformats.org/presentationml/2006/main}}{kind}/a:lvl1pPr/a:defRPr", NS)
        if node is not None and node.get("sz"):
            out[kind] = int(node.get("sz")) / 100
    return out


def layouts(prs):
    found = []
    for mi, master in enumerate(prs.slide_masters):
        for lay in master.slide_layouts:
            phs = []
            for ph in lay.placeholders:
                phs.append({
                    "type": str(ph.placeholder_format.type).split(".")[-1].split(" ")[0],
                    "x": pt(ph.left or 0), "y": pt(ph.top or 0),
                    "w": pt(ph.width or 0), "h": pt(ph.height or 0),
                })
            pics = sum(1 for s in lay.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
            found.append({"master": mi, "name": lay.name, "placeholders": phs, "pictures": pics})
    return found


def guess_roles(lays):
    roles = {}
    for role, hints in ROLE_HINTS.items():
        # hint order is preference order: "title only" beats "title and content"
        match = next((lay["name"] for h in hints for lay in lays if h in lay["name"].lower()), None)
        if match:
            roles[role] = match
    return roles


def usage(prs):
    sizes, colors, fonts = collections.Counter(), collections.Counter(), collections.Counter()
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    if not r.text.strip():
                        continue
                    n = len(r.text)
                    if r.font.size:
                        sizes[r.font.size.pt] += n
                    try:
                        if r.font.color and r.font.color.type is not None and r.font.color.rgb:
                            colors[str(r.font.color.rgb)] += n
                    except AttributeError:
                        pass  # theme colors without an explicit RGB
                    if r.font.name:
                        fonts[r.font.name] += n
    return sizes, colors, fonts


def draft(src: Path, out: Path, report: Path | None):
    prs = Presentation(str(src))
    base = yaml.safe_load(DEFAULT_BRAND.read_text(encoding="utf-8"))
    W, H = pt(prs.slide_width), pt(prs.slide_height)
    colors, fonts = theme(prs)
    lays = layouts(prs)
    roles = guess_roles(lays)
    sizes, used_colors, used_fonts = usage(prs)
    notes = []

    b = base
    b["name"] = f"Draft from {src.stem}"
    b["canvas"] = {"width": W, "height": H}
    sx, sy = W / 960, H / 540  # scale default geometry to the template canvas
    for key in ("margin_left", "margin_right", "gutter", "rule_w"):
        b["layout"][key] = round(b["layout"][key] * sx, 1)
    for key in ("eyebrow_y", "title_y", "title_h", "rule_y", "body_top", "body_bottom", "footer_y"):
        b["layout"][key] = round(b["layout"][key] * sy, 1)

    # geometry from the content layout's title and body placeholders
    content = next((l for l in lays if l["name"] == roles.get("content")), None)
    body_layout = next((l for l in lays if "content" in l["name"].lower() or "내용" in l["name"]), None)
    if content:
        title = next((p for p in content["placeholders"] if p["type"] in ("TITLE", "CENTER_TITLE")), None)
        if title:
            b["layout"]["margin_left"] = title["x"]
            b["layout"]["margin_right"] = round(W - title["x"] - title["w"], 1)
            b["layout"]["title_y"] = title["y"]
            b["layout"]["title_h"] = title["h"]
            b["layout"]["rule_y"] = None  # the template draws its own title treatment
            b["layout"]["eyebrow_y"] = max(8.0, round(title["y"] - 18, 1))
            b["layout"]["body_top"] = round(title["y"] + title["h"] + 18, 1)
            notes.append(f"title band taken from layout '{content['name']}'")
    if body_layout:
        body = next((p for p in body_layout["placeholders"] if p["type"] in ("BODY", "OBJECT")), None)
        if body:
            b["layout"]["body_top"] = body["y"]
            b["layout"]["body_bottom"] = round(body["y"] + body["h"], 1)
            notes.append(f"body area taken from layout '{body_layout['name']}'")
    footer = next((p for l in lays for p in l["placeholders"] if p["type"] in ("FOOTER", "SLIDE_NUMBER")), None)
    if footer:
        b["layout"]["footer_y"] = footer["y"]

    # colors: theme slots → brand roles
    pick = lambda *keys: next((colors[k] for k in keys if colors.get(k)), None)
    dark, light = pick("dk1"), pick("lt1")
    mapping = {"ink": dark, "body": pick("dk2", "dk1"), "surface": light, "surface_alt": pick("lt2"),
               "accent": pick("accent1"), "accent_2": pick("accent2"), "dark": pick("dk2", "dk1")}
    for role, value in mapping.items():
        if value:
            b["colors"][role] = value.upper()
    series = [colors[k].upper() for k in ("accent1", "accent2", "accent3", "accent4", "accent5", "accent6")
              if colors.get(k)]
    if series:
        b["colors"]["series"] = series[:5]
    if b["colors"].get("accent") and b["colors"].get("surface"):
        ratio = contrast(b["colors"]["accent"], b["colors"]["surface"])
        if ratio < 4.5:
            notes.append(f"accent {b['colors']['accent']} has {ratio:.1f}:1 contrast on the background; "
                         "use it for shapes and rules, not small text")
    if b["colors"].get("dark") and luminance(b["colors"]["dark"]) > 0.4:
        b["colors"]["dark"] = base["colors"]["ink"]
        notes.append("theme has no dark color; `dark` kept at the ink color — check covers")

    ms = master_sizes(prs)
    if ms.get("titleStyle"):
        b["type"]["title"] = ms["titleStyle"]
        notes.append(f"type.title set to the master title size {ms['titleStyle']}pt; headlines must fit one "
                     "line at that size — if they rarely do, agree a smaller title size in the template")
    if ms.get("bodyStyle"):
        notes.append(f"master body size is {ms['bodyStyle']}pt; type.lead / type.body kept at the defaults — "
                     "align them with your template if they should match")

    major, minor = fonts.get("majorFont", {}), fonts.get("minorFont", {})
    if major.get("latin"):
        b["fonts"]["heading"]["latin"] = major["latin"]
    if major.get("east_asian"):
        b["fonts"]["heading"]["east_asian"] = major["east_asian"]
    if minor.get("latin"):
        b["fonts"]["body"]["latin"] = minor["latin"]
    if minor.get("east_asian"):
        b["fonts"]["body"]["east_asian"] = minor["east_asian"]

    # template mode
    rel = Path("assets") / src.name
    target = out.parent / rel
    if src.resolve() != target.resolve():
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, target)
    b["template"]["path"] = str(rel)
    for role in ("cover", "section", "content", "closing", "blank"):
        b["template"]["layouts"][role] = roles.get(role)
    b["backgrounds"] = {k: {"color": v, "image": None} for k, v in
                        (("cover", "surface"), ("section", "surface"), ("closing", "surface"),
                         ("content", "surface"))}
    notes.append("backgrounds and logos come from the template masters; set "
                 "template.draw_backgrounds / template.draw_chrome to true to draw the brand's own")

    if sizes:
        common = [s for s, _ in sizes.most_common(6)]
        notes.append(f"most used font sizes in the slides (pt): {', '.join(str(s) for s in common)}")
    if used_colors:
        notes.append("most used text colors: " + ", ".join(c for c, _ in used_colors.most_common(6)))
    if used_fonts:
        notes.append("fonts set directly on text: " + ", ".join(f for f, _ in used_fonts.most_common(4)))

    out.parent.mkdir(parents=True, exist_ok=True)
    header = (f"# Draft brand profile extracted from {src.name}.\n"
              "# Review every value; decisions that are not numbers go to guidelines.md.\n")
    out.write_text(header + yaml.safe_dump(b, sort_keys=False, allow_unicode=True), encoding="utf-8")

    lines = [f"# Extraction report — {src.name}", "", f"Canvas: {W} x {H} pt", "",
             "## Theme colors", ""]
    lines += [f"- {k}: {v}" for k, v in colors.items()]
    lines += ["", "## Theme fonts", "", f"- headings: {major}", f"- body: {minor}", "",
              "## Layouts", "", "| master | layout | placeholders | pictures |", "|---|---|---|---|"]
    for l in lays:
        phs = ", ".join(f"{p['type']}@({p['x']},{p['y']} {p['w']}x{p['h']})" for p in l["placeholders"])
        lines.append(f"| {l['master']} | {l['name']} | {phs} | {l['pictures']} |")
    lines += ["", "## Role mapping (guessed)", ""]
    lines += [f"- {r}: {roles.get(r) or 'not found — set template.layouts.' + r}" for r in ROLE_HINTS]
    lines += ["", "## Notes", ""] + [f"- {n}" for n in notes]
    text = "\n".join(lines) + "\n"
    if report:
        report.write_text(text, encoding="utf-8")
    return text


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx", help="company template (.pptx/.potx saved as .pptx) or a representative deck")
    ap.add_argument("-o", "--output", default="brand/brand.draft.yaml")
    ap.add_argument("--report", help="write the extraction report (markdown) here")
    args = ap.parse_args()
    out = Path(args.output)
    report = draft(Path(args.pptx), out, Path(args.report) if args.report else None)
    print(report)
    print(f"draft written to {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
